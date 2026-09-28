"""Read exported workbook and independently reconstruct the linked baseline."""
import json,math,bisect
from pathlib import Path
from statistics import NormalDist
import openpyxl
P=Path(__file__).parent; ROOT=P.parents[1]
d=json.loads((P/'inputs.json').read_text())
file=ROOT/'outputs/cprt-linked-20260928/CPRT_Linked_Service_Revenue.xlsx'
w=openpyxl.load_workbook(file,data_only=True);f=openpyxl.load_workbook(file,data_only=False)
def eq(a,b):
 assert isinstance(a,(int,float)) and math.isclose(a,b,rel_tol=2e-8,abs_tol=1e-7),(a,b)
errors=[(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e'];assert not errors,errors
assert not f._external_links
for name in ['Source Data','Fee Data','Evidence']:
 assert not any(c.data_type=='f' for row in f[name] for c in row),name
assert w.sheetnames[:4]==['RPM Summary','Volume Build','Vehicle Economics','Revenue Bridge']
N=NormalDist();quarters=[];node_checks=0
std=[r for r in d['fees'] if r['page']=='non-licensed' and r['title_group']=='non-clean' and r['vehicle_class']=='standard' and r['fee_type']=='buyer_fee' and r['payment_method']=='secured']
pref=[r for r in d['fees'] if 'high' in r['page'] and r['title_group']=='non-clean' and r['vehicle_class']=='standard' and r['fee_type']=='buyer_fee' and r['payment_method']=='secured']
virtual=[r for r in d['fees'] if r['page']=='non-licensed' and r['title_group']=='non-clean' and r['vehicle_class']=='standard' and r['fee_type']=='virtual_bid_pre_bid']
def fee(price,grid):
 r=grid[bisect.bisect_right([float(x['band_low_usd']) for x in grid],price)-1]
 v=virtual[bisect.bisect_right([float(x['band_low_usd']) for x in virtual],price)-1]
 return float(r['fee_usd'] or 0)+price*float(r['fee_pct'] or 0)/100+float(v['fee_usd'])+110
for p,period in enumerate(d['periods']):
 weights={}
 for b in range(4):
  for k in range(6):
   weights[b,k]=sum(sum(birth*tw for birth,tw in zip(r['births'],period['weights']))*r['survival']*d['split'][b]*r['relative_claim_weight'] for r in d['fleet'] if r['body']==b and r['bucket']==k)
   eq(w['Fleet Engine'].cell(194+b*6+k,6+p).value,weights[b,k])
 share=sum(c['weights'][p]*c['allocations'][p] for c in d['carriers'])/sum(c['weights'][p] for c in d['carriers'])
 eq(w['Carrier Engine'].cell(91+p,4).value,share)
 tl=core=proceeds=0
 for b in range(4):
  for k in range(6):
   acv=10000*math.exp(-.1*(d['calibration'][k]['representative_age']-10))*d['value_ratios'][b]
   mu=d['calibration'][k]['log_repair_median']+math.log(d['repair_ratios'][b])
   for j in range(9):
    recovery=.3+.2*(.5-(j+.5)/9)
    threshold=acv*(1-recovery*.96)
    cut=N.cdf((math.log(threshold)-mu)/d['sigma'])
    prob=max(0,(j+1)/9-max(j/9,cut));price=acv*recovery
    rate=(fee(price,std)+fee(price,pref))/2+.04*price
    units=weights[b,k]*prob;tl+=units;core+=units*rate;proceeds+=units*price
    row=6+p*216+(b*6+k)*9+j
    for c,val in [(14,prob),(15,price),(20,rate),(21,units)]:eq(w['Damage Engine'].cell(row,c).value,val)
    node_checks+=1
 rpu=core/tl+55;asp=proceeds/tl;rawunits=tl*share
 eq(w['Vehicle Economics'].cell(22,4+p).value,rpu)
 eq(w['Vehicle Economics'].cell(9,4+p).value,asp)
 quarters.append({'quarter':period['label'],'rawunits':rawunits,'rpu':rpu,'asp':asp,'share':share})
scale=d['actuals'][3]['us_service']*1e6*.9/(quarters[3]['rawunits']*quarters[3]['rpu'])
for p,q in enumerate(quarters):
 q['units']=q['rawunits']*scale;q['insurance_m']=q['units']*q['rpu']/1e6;q['other_us_m']=d['actuals'][3]['us_service']*.1;q['intl_m']=d['actuals'][3]['intl_service'];q['total_m']=sum(q[k] for k in ['insurance_m','other_us_m','intl_m'])
 eq(w['Volume Build'].cell(23,4+p).value,q['units']);eq(w['Revenue Bridge'].cell(21,4+p).value,q['total_m'])
 if p>=4:eq(w['RPM Summary'].cell(12,4+p).value,q['total_m'])
for r in [8,9,10,11,12]:
 for c in range(4,12):eq(w['Checks'].cell(r,c).value,0)
prop=json.loads((P/'propagation_tests.json').read_text());assert prop['actualsUnchanged'] and prop['before']==prop['restored'];assert prop['before']!=prop['adoption'] and prop['before']!=prop['repair']
assert w['RPM Summary']['H28'].value=='n.a.'
out={'passed':True,'damage_nodes_independently_checked':node_checks,'formula_errors':errors,'data_tabs_values_only':True,'external_links':False,'quarter_results':quarters,'historical_us_residual_m':[w['Checks'].cell(16,c).value for c in range(4,8)],'calibration':'One FY26Q4 dollar anchor; not observed absolute unit levels','engine':'Artifact Tool recalculation + independent Python; desktop Excel not tested'}
(P/'validation.json').write_text(json.dumps(out,indent=2));print(json.dumps({k:v for k,v in out.items() if k!='quarter_results'},indent=2))
