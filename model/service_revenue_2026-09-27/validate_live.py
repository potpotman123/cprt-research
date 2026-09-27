"""Independent numerical checks; openpyxl is used only to read the exported file."""
import json, math, bisect, os
from pathlib import Path
import openpyxl
p=Path(__file__).parent
out=Path(os.environ.get('CPRT_OUTPUT_DIR',p.parent if p.name=='work' else p))
d=json.loads((p/'model_inputs.json').read_text())
w=openpyxl.load_workbook(out/'CPRT_Fleet_and_Body_Mix.xlsx',data_only=True)
f=openpyxl.load_workbook(out/'CPRT_Fleet_and_Body_Mix.xlsx',data_only=False)
def eq(actual,expected):
    assert isinstance(actual,(float,int)) and math.isclose(actual,expected,rel_tol=2e-8,abs_tol=1e-7),(actual,expected)
errors=[(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e']
assert not errors,errors
assert not f._external_links
assert len(w.sheetnames)==14
assert w.sheetnames[:4]==['RPM Summary','Volume Build','Vehicle Economics','Revenue Bridge']
assert [s.title for s in f if s.sheet_properties.tabColor]==['RPM Summary','CALCULATIONS','DATA']
assert not any(c.data_type=='f' for s in f if s.title.startswith('AD ') for row in s for c in row)
assert f['RPM Summary']['C8'].font.name=='Garamond'
mix=[next(r for r in d['crss'] if r['body']==b and r['age']=='7-9') for b in ['Car','SUV','Pickup','Minivan']]
weights=[(float(r['front_share']),float(r['rear_share']),1-float(r['front_share'])-float(r['rear_share'])) for r in mix]
scale=3682/sum(v*k for v,k in zip(weights[0],[1,.6,.8]))
aaa=d['aaa']['rows']; repairs=[]
for b,key in enumerate(['camry','rogue','f150','rogue']):
    costs=[scale*aaa[0][key]/aaa[0]['camry'],scale*.6*aaa[2][key]/aaa[2]['camry'],scale*.8*[1,1.1,1.15,1.1][b]]
    repairs.append(sum(v*k for v,k in zip(weights[b],costs)))
    eq(w['Repair Costs'].cell(25,4+b).value,repairs[-1])
sales={int(r['year']):(float(r['cars_k']),float(r['light_trucks_k'])) for r in d['sales']}
sales[2026]=sales[2025];sales[2027]=sales[2025]
split=[.409/.565,.109/.565,.047/.565]
def survival(a,b):
    z=a/(1.270 if b==0 else 1.099)
    if z>=31:return 0.
    lo=math.floor(z);k='survival_cars' if b==0 else 'survival_light_trucks'
    return float(d['survival'][lo][k])+(z-lo)*(float(d['survival'][lo+1][k])-float(d['survival'][lo][k]))
def exposure(a):return (.5 if a==0 else 1)*math.exp(-.09027*max(a-6,0))
def bucket(a):return 0 if a==0 else 1 if a<=3 else 2 if a<=6 else 3 if a<=9 else 4 if a<=12 else 5
stock={(y,b,a):(sales[y-a][0] if b==0 else sales[y-a][1]*split[b-1])*survival(a,b) for y in range(2023,2028) for b in range(4) for a in range(46)}
offset=[math.log((1 if b==0 else 1.5)/(repairs[b]/repairs[0]))/d['repair']['inherited_sigma'] for b in range(4)]
cdf=lambda x:(1+math.erf(x/math.sqrt(2)))/2
prob={}
for k in range(6):
    wt=[sum(stock[2025,b,a]*exposure(a) for a in range(46) if bucket(a)==k) for b in range(4)]
    lo,hi=-8,8;target=float(d['ccc'][k]['cy2025'])/100
    for _ in range(60):
        z=(lo+hi)/2;avg=sum(v*cdf(z-offset[b]) for b,v in enumerate(wt))/sum(wt)
        if avg<target:lo=z
        else:hi=z
    for b in range(4):
        prob[k,b]=cdf(z-offset[b]);eq(w['Loss Calibration'].cell(8+k,5+b).value,prob[k,b])
fees=[r for r in d['fees'] if r['page']=='non-licensed' and r['title_group']=='non-clean' and r['vehicle_class']=='standard' and r['fee_type']=='buyer_fee' and r['payment_method']=='secured']
virtual=[r for r in d['fees'] if r['page']=='non-licensed' and r['title_group']=='non-clean' and r['vehicle_class']=='standard' and r['fee_type']=='virtual_bid_pre_bid']
def buyer(x):
    band=fees[bisect.bisect_right([float(r['band_low_usd']) for r in fees],x)-1]
    vb=virtual[bisect.bisect_right([float(r['band_low_usd']) for r in virtual],x)-1]
    return float(band['fee_usd'] or 0)+x*float(band['fee_pct'] or 0)/100+float(vb['fee_usd'])+110
def rpu(a,b):
    price=10000*math.exp(-.1*(a-10))*(1 if b==0 else 1.5)*.3
    return sum(wt*buyer(price*m) for wt,m in zip([.25,.5,.25],[.5,1,1.5]))+.04*price
annual=[]
for y in range(2023,2028):
    totals=[0.,0.,0.,0.]
    for b in range(4):
        for a in range(46):
            s=stock[y,b,a];cl=s*exposure(a);u=cl*prob[bucket(a),b];v=u*rpu(a,b)
            eq(w['Cohort Roll'].cell(42+b*46+a,4+y-2023).value,s)
            eq(w['Fee Calculation'].cell(8+b*46+a,13).value,rpu(a,b))
            totals=[old+new for old,new in zip(totals,[s,cl,u,v])]
    annual.append(totals)
    for row,val in zip([30,31,32,33],totals):eq(w['Cohort Roll'].cell(row,4+y-2023).value,val)
quarters=[]
for j,period in enumerate(d['periods']):
    eq(sum(period['weights']),1)
    totals=[sum(a[k]*wt for a,wt in zip(annual,period['weights'])) for k in range(4)]
    quarters.append(totals)
    for row,val in zip([8,9,10,11],totals):eq(w['Cohort Roll'].cell(row,4+j).value,val)
growth=[None]*4+[quarters[j][3]/quarters[j-4][3]-1 for j in range(4,12)]
ug=d['actuals'][7]['us_service']/d['actuals'][3]['us_service']-1
ig=d['actuals'][7]['intl_service']/d['actuals'][3]['intl_service']-1
forecasts=[]
for j in range(4):
    old=d['actuals'][4+j];baseline=old['us_service']*(1+ug)
    delta=baseline/(1+.9*growth[7])*.9*(growth[8+j]-growth[7])
    us=baseline+delta;intl=old['intl_service']*(1+ig);total=us+intl
    for row,val in [(8,us),(9,intl),(10,total)]:eq(w['RPM Summary'].cell(row,8+j).value,val)
    eq(w['Revenue Bridge'].cell(17,8+j).value,delta)
    eq(w['RPM Summary'].cell(10,4+j).value,old['us_service']+old['intl_service'])
    forecasts.append({'quarter':d['periods'][8+j]['label'],'total_service_m':total,'fleet_adjustment_m':delta})
eq(w['RPM Summary']['L10'].value,sum(r['total_service_m'] for r in forecasts[:2]))
assert w['Revenue Bridge']['D34'].value<1e-7
tests=json.loads((p/'live_model_tests.json').read_text())
assert tests['before']==tests['restored'] and tests['before']!=tests['changed'] and tests['before']!=tests['repairChanged'] and tests['zeroExposure']==0
result={'passed':True,'formula_errors':errors,'external_links':False,'data_tabs_values_only':True,'colored_section_markers_only':True,'independent_stock_cells_checked':920,'repair_means':repairs,'forecasts':forecasts,'input_propagation_tests_passed':True,'engine':'artifact-tool recalculation plus independent Python reconstruction; not desktop Excel tested'}
(p/'live_validation.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2))
