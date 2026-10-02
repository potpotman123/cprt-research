"""Edits to 'model/CPRT Model V4.xlsx' (owner's file of record).
1. Reverse DCF: FY27–28 revenue, EBITDA and D&A read D Barclays (11 Sep 2026) exactly; H5 = Barclays revenue CAGR; exit multiple solved live
   so the implied price equals the current price on Barclays' path.
2. Summary rebuilt in the Solstice layout: stacked tables (headline, Street, delta to Street), component delta by thesis, EBITDA bridge,
   valuation strip; no blank rows; dashes where no line exists."""
import sys, pathlib; sys.path.insert(0, str(pathlib.Path(__file__).parent))
import openpyxl
from cprt_model.common import *
P = ROOT / 'model/CPRT Model V4.xlsx'; wb = openpyxl.load_workbook(P); W = 'WACC, Tax, and Reverse DCF'; ws = wb[W]
# ------------------------------------------------------------------ 1. reverse DCF on Barclays' path
put(ws, 'B5', 'Barclays revenue CAGR, FY26–FY28 (D Barclays, 11 Sep 2026; the FY27–28 path below is Barclays\' exactly)', 'label'); put(ws, 'H5', '=(J10/H10)^(1/2)-1', 'formula', F_PCT2, b=True)
put(ws, 'I10', "='D Barclays'!E5", 'link', F_MONEY, b=True); put(ws, 'J10', "='D Barclays'!F5", 'link', F_MONEY, b=True)
put(ws, 'I13', "='D Barclays'!E6", 'link', F_MONEY, b=True); put(ws, 'J13', "='D Barclays'!F6", 'link', F_MONEY, b=True)
put(ws, 'I14', '=I13/I10', 'formula', F_PCT); put(ws, 'J14', '=J13/J10', 'formula', F_PCT)
put(ws, 'I17', "='D Barclays'!E6-'D Barclays'!E7", 'link', F_MONEY); put(ws, 'J17', "='D Barclays'!F6-'D Barclays'!F7", 'link', F_MONEY)
put(ws, 'B36', "Exit EV / EBITDA multiple — solved so the implied price (H51) equals the current price (H52) on Barclays' path", 'label')
put(ws, 'H36', '=(H52*H50-SUM(H45:H48)-H40)/(H37*J31)', 'formula', F_X, b=True)
put(ws, 'R36', 'Required EV (price × FDSO − bridge) less PV of FY27–28 UFCF, divided by FY28E EBITDA × its discount factor; no Goal Seek needed. Barclays\' own PT uses 10× on FY26', 'note')
put(ws, 'B56', 'Barclays revenue CAGR, FY26–FY28', 'label')
# ------------------------------------------------------------------ 2. Summary (Solstice layout)
idx = wb.sheetnames.index('Summary'); wb.remove(wb['Summary']); sm = wb.create_sheet('Summary', idx); tab_color(sm, NAVY)
setup(sm, label_w=46, ncols=6, notes_col='L'); NOTES_COL['Summary'] = 'L'
for c in 'DEFGHI': sm.column_dimensions[c].width = 11.5
title(sm, 'Copart, Inc. (NASDAQ: CPRT) — Summary', 'Model = 3SM on the selected RPM case. Street = Barclays, 11 Sep 2026 (Underweight, $25). Δ = model − Street. $M except per share; fiscal years end 31 July. – = no line.')
header(sm, 4, None, hist_span=(4, 6), proj_span=(7, 9), notes_col='L')
YEARS = list(range(2024, 2030)); SC = {y: L(4 + i) for i, y in enumerate(YEARS)}; TC = {y: L(4 + y - 2022) for y in YEARS}; BC = {2026: 'D', 2027: 'E', 2028: 'F', 2029: 'G'}
per = [f'FY{y}{"A" if y <= 2026 else "E"}' for y in YEARS]; T = {'rev': 8, 'ebitda': 18, 'eps': 28, 'fcf': 112}
R = {}; cur = [4]
def nxt(): cur[0] += 1; return cur[0]
def row(key, text, units, fn, fmt=F_MONEY, kind='link', bold=False, ital=False, note='', years=YEARS):
    r = nxt(); R[key] = r; label(sm, r, text, units, b=bold, i=ital, note=note)
    for fy in YEARS:
        c = SC[fy]; v = fn(fy, c, SC.get(fy - 1)) if fy in years else None
        if v is None: put(sm, f'{c}{r}', DASH, 'note').alignment = Alignment(horizontal='center')
        else: put(sm, f'{c}{r}', v, kind, fmt, b=bold)
    return r
g = lambda k: (lambda fy, c, p: f'=IF({p}{R[k]}=0,0,{c}{R[k]}/{p}{R[k]}-1)' if p else None)
m = lambda n, d: (lambda fy, c, p: f'=IF({c}{R[d]}=0,0,{c}{R[n]}/{c}{R[d]})')
tsm = lambda k: (lambda fy, c, p: f"='3SM'!{TC[fy]}{T[k]}")
bar = lambda r_: (lambda fy, c, p: f"='D Barclays'!{BC[fy]}{r_}")
section(sm, nxt(), 'Headline financials — model (3SM, selected case on RPM D3)', per)
row('rev', 'Revenue', '$M', tsm('rev'), bold=True, note='3SM row 8 ← RPM')
row('revg', '% yoy growth', '%', g('rev'), F_PCT, 'formula', ital=True, years=YEARS[1:])
row('ebitda', 'Adj. EBITDA', '$M', tsm('ebitda'), bold=True, note='3SM row 18: EBIT + D&A (Barclays definition)')
row('ebitdam', '% margin', '%', m('ebitda', 'rev'), F_PCT, 'formula', ital=True)
row('eps', 'GAAP diluted EPS', '$', tsm('eps'), F_USD + '0', bold=True, note='3SM row 28')
row('epsg', '% yoy growth', '%', g('eps'), F_PCT, 'formula', ital=True, years=YEARS[1:])
row('fcf', 'Free cash flow (CFO + capex)', '$M', tsm('fcf'), bold=True, note='3SM row 112')
row('fcfm', '% margin', '%', m('fcf', 'rev'), F_PCT, 'formula', ital=True)
row('ufcf', 'Unlevered FCF (DCF)', '$M', lambda fy, c, p: f'=DCF!{TC[fy]}59', note='DCF row 59')
row('ufcfm', '% margin', '%', m('ufcf', 'rev'), F_PCT, 'formula', ital=True)
section(sm, nxt(), 'Street — Barclays estimates (11 Sep 2026)', per)
row('brev', 'Revenue', '$M', bar(5), years=list(BC), note='D Barclays row 5 (numeric extraction; licensed report held locally)')
row('brevg', '% yoy growth', '%', g('brev'), F_PCT, 'formula', ital=True, years=[2027, 2028, 2029])
row('beb', 'Adj. EBITDA', '$M', bar(6), years=list(BC), note='D Barclays row 6')
row('bebm', '% margin', '%', m('beb', 'brev'), F_PCT, 'formula', ital=True, years=list(BC))
row('beps', 'Diluted EPS', '$', bar(10), F_USD + '0', years=list(BC), note='D Barclays row 10')
row('bepsg', '% yoy growth', '%', g('beps'), F_PCT, 'formula', ital=True, years=[2027, 2028, 2029])
section(sm, nxt(), 'Delta to Street — model less Barclays', per)
row('drev', 'Revenue Δ', '$M', lambda fy, c, p: f'={c}{R["rev"]}-{c}{R["brev"]}', kind='formula', years=list(BC))
row('drevp', 'Revenue Δ %', '%', lambda fy, c, p: f'=IF({c}{R["brev"]}=0,0,{c}{R["rev"]}/{c}{R["brev"]}-1)', F_PCT, 'formula', bold=True, years=list(BC))
row('deb', 'Adj. EBITDA Δ', '$M', lambda fy, c, p: f'={c}{R["ebitda"]}-{c}{R["beb"]}', kind='formula', years=list(BC))
row('debp', 'Adj. EBITDA Δ %', '%', lambda fy, c, p: f'=IF({c}{R["beb"]}=0,0,{c}{R["ebitda"]}/{c}{R["beb"]}-1)', F_PCT, 'formula', bold=True, years=list(BC))
row('deps', 'EPS Δ', '$', lambda fy, c, p: f'={c}{R["eps"]}-{c}{R["beps"]}', F_USD + '0', 'formula', years=list(BC))
row('depsp', 'EPS Δ %', '%', lambda fy, c, p: f'=IF({c}{R["beps"]}=0,0,{c}{R["eps"]}/{c}{R["beps"]}-1)', F_PCT, 'formula', bold=True, years=list(BC))
br = lambda k: 7 + (k - 1) * 13 + 1 + 8; case = lambda k, fy: f"Scenarios!{'Q' if fy == 2027 else 'R'}{br(k)}"
section(sm, nxt(), 'Component delta to Street — legacy service revenue by thesis ($M; engines run to FY28)', per)
row('c0', 'Street-implied (case 0)', '$M', lambda fy, c, p: f'={case(1, fy)}', years=[2027, 2028], note=f'Scenarios row {br(1)}; FY27 reverse-solved to JPMorgan $4,061m (D Facts)')
row('t1', 'Thesis 1 — coverage persistence', '$M', lambda fy, c, p: f'={case(2, fy)}-{case(1, fy)}', bold=True, years=[2027, 2028], note=f'case 1 − case 0 (Scenarios rows {br(2)} − {br(1)})')
row('t2', 'Thesis 2 — aftermarket substitution', '$M', lambda fy, c, p: f'={case(4, fy)}-{case(2, fy)}', bold=True, years=[2027, 2028], note=f'thesis 2 case − case 1 (rows {br(4)} − {br(2)})')
row('t3', 'Thesis 3 — carrier reallocation', '$M', lambda fy, c, p: f'={case(5, fy)}-{case(2, fy)}', bold=True, years=[2027, 2028], note=f'thesis 3 case − case 1 (rows {br(5)} − {br(2)})')
row('inter', 'Interaction', '$M', lambda fy, c, p: f'={case(6, fy)}-{case(1, fy)}-{c}{R["t1"]}-{c}{R["t2"]}-{c}{R["t3"]}', kind='formula', ital=True, years=[2027, 2028], note='all three − case 0 − sum of parts (common base)')
row('all3', 'All three vs Street-implied', '$M', lambda fy, c, p: f'={case(6, fy)}-{case(1, fy)}', bold=True, years=[2027, 2028])
row('sel', 'Selected case vs Street-implied', '$M', lambda fy, c, p: f"=RPM!{'I' if fy == 2027 else 'J'}10-{c}{R['c0']}", kind='link', bold=True, years=[2027, 2028], note='RPM I10/J10 less case 0')
row('veh', 'Vehicle sales vs JPMorgan ($710m FY27E)', '$M', lambda fy, c, p: '=RPM!I12-Street!F8', years=[2027], note='Street F8')
section(sm, nxt(), 'EBITDA bridge to Barclays ($M)', per)
row('b_drev', 'Revenue Δ', '$M', lambda fy, c, p: f'={c}{R["drev"]}', kind='formula', years=[2027, 2028, 2029])
row('b_reff', "of which revenue effect at Barclays' margin", '$M', lambda fy, c, p: f'={c}{R["drev"]}*{c}{R["bebm"]}', kind='formula', ital=True, years=[2027, 2028, 2029])
row('b_meff', 'of which margin effect (facility cost per unit, G&A ratio)', '$M', lambda fy, c, p: f'={c}{R["deb"]}-{c}{R["b_reff"]}', kind='formula', ital=True, years=[2027, 2028, 2029])
row('b_deb', 'Adj. EBITDA Δ', '$M', lambda fy, c, p: f'={c}{R["deb"]}', kind='formula', bold=True, years=[2027, 2028, 2029])
section(sm, nxt(), 'Valuation', ['Price (Cover)', 'DCF implied', 'Upside', 'WACC', 'Exit multiple', 'Tax rate'])
r = nxt(); label(sm, r, 'Cover D6 · DCF H98 · DCF H99 · WACC tab N31 · DCF H84 (perpetuity-implied) · WACC tab N39')
for c, f_, fmt in (('D', '=Cover!D6', F_USD), ('E', '=DCF!H98', F_USD), ('F', '=DCF!H99', F_PCT), ('G', f"='{W}'!N31", F_PCT2), ('H', '=DCF!H84', F_X), ('I', f"='{W}'!N39", F_PCT2)): put(sm, f'{c}{r}', f_, 'link', fmt, b=(c == 'E'))
r = nxt(); label(sm, r, 'Reverse DCF: exit multiple the current price implies on Barclays\' FY27–28 path (WACC tab H36)'); put(sm, f'H{r}', f"='{W}'!H36", 'link', F_X, b=True)
for c in 'DEFGI': put(sm, f'{c}{r}', DASH, 'note').alignment = Alignment(horizontal='center')
sm.freeze_panes = 'D5'; wb.save(P); print('saved; Summary rows:', cur[0])
