"""V5: (1) DCF row 85 — reference multiples beside the implied exit multiple (reverse-DCF market-implied, Barclays PT 10x, Barclays-cited
trading 12.7x and 10-yr average 16.0x, the last three recorded on D Facts with sources); (2) RPM reformat — freeze panes at D6 (was D44),
notes moved from P to O beside the data, long labels shortened, sub-headers as group bars, dashes in empty data cells, list column hidden."""
import sys, pathlib; sys.path.insert(0, str(pathlib.Path(__file__).parent))
import openpyxl
from cprt_model.common import *
P = pathlib.Path(sys.argv[1]); wb = openpyxl.load_workbook(P); W = 'WACC, Tax, and Reverse DCF'
# ---------------------------------------------------------------- D Facts: Barclays multiples with sources
df = wb['D Facts']; facts = [(54, 'Barclays price-target multiple: 10× EV/EBITDA applied to 2026 forecast ($25 PT)', 'barclays_pt_multiple', 10, 'Valuation methodology paragraph, Barclays Copart F4Q26 Review, 11 Sep 2026 (also 25 Aug 2026 note)'),
                             (55, 'Copart EV/EBITDA on forward 12-month estimates, as cited by Barclays', 'barclays_cited_fwd_multiple', 12.7, 'Barclays, CPRT vs. RBA Salvage Auction Wars, 25 Aug 2026, valuation paragraph'),
                             (56, 'Copart 10-year average forward EV/EBITDA, as cited by Barclays', 'barclays_cited_10y_avg_multiple', 16.0, 'Same paragraph (25 Aug 2026 note)')]
for r, lab, key, val, srcnote in facts:
    put(df, f'B{r}', lab, 'label'); put(df, f'C{r}', key, 'label', i=True); put(df, f'D{r}', val, 'input', F_X); put(df, f'E{r}', 'x', 'label', i=True); put(df, f'F{r}', 'VERIFIED (licensed report held locally; number only)', 'note'); put(df, f'G{r}', srcnote, 'note'); put(df, f'H{r}', '2026-09-11' if r == 54 else '2026-08-25', 'note'); put(df, f'I{r}', 'research/cprt_discovery_plan_2026-09-26/sources (local; not committed)', 'note')
# ---------------------------------------------------------------- DCF row 85: reference multiples
d = wb['DCF']
put(d, 'B85', "   Reference multiples (x): H market-implied on Barclays' path (reverse DCF) · I Barclays PT, 10× FY26 EBITDA · J Barclays-cited trading, 12.7× forward · K 10-yr average, 16.0×", 'label', i=True)
put(d, 'H85', f"='{W}'!H36", 'link', F_X, b=True); put(d, 'I85', "='D Facts'!D54", 'link', F_X); put(d, 'J85', "='D Facts'!D55", 'link', F_X); put(d, 'K85', "='D Facts'!D56", 'link', F_X)
put(d, 'P85', 'Sources: reverse DCF H36 (solved so the implied price equals the current price on Barclays\' FY27–28 path); D Facts rows 54–56 (Barclays notes, 11 Sep and 25 Aug 2026)', 'note')
put(d, 'P80', 'Terminal value ÷ FY2031E EBITDA. Compare with the reference multiples in row 85', 'note')
# ---------------------------------------------------------------- RPM reformat
r_ = wb['RPM']
r_.freeze_panes = 'D6'
for col, w in (('B', 52), ('C', 7), ('N', 12.5), ('O', 64), ('P', 9)): r_.column_dimensions[col].width = w
for c in 'DEFGHIJKLM': r_.column_dimensions[c].width = 11.5
r_.column_dimensions['U'].hidden = True
for row in range(1, r_.max_row + 1):                                   # notes P → O (adjacent to the CAGR column)
    v = r_[f'P{row}'].value
    if v is not None: r_[f'O{row}'].value = v; r_[f'O{row}'].font = font(C_NOTE, i=True, sz=10); r_[f'P{row}'].value = None
NOTES_COL['RPM'] = 'O'; put(r_, 'O4', 'Notes / sources', 'note')
short = {'B3': 'Scenario selector — click D3 to choose the case (index in I3)', 'B42': 'Beyond the engine horizon (FY29–31): fade to long-run anchors (rows 53–60)',
         'B43': 'Service revenue growth FY29–31 — derived; type a rate in D43 to override', 'B45': 'US insurance unit growth FY29–31 — derived per year (row 21); anchor shown',
         'B47': 'Δ to Street and to the engine (FY27E legacy service revenue, $M)', 'B53': 'FY29–31 fade: g(t) = anchor + (g_FY28 − anchor) × decay^(t−2028)',
         'B54': 'Long-run fleet growth, CY2027→CY2028 (E1 fleet roll)', 'B55': 'Total-loss-frequency drift beyond FY28 (pts a year)', 'B56': 'US insurance unit growth — long-run anchor',
         'B57': 'Used-vehicle price inflation anchor (BLS used-car CPI, 10-yr CAGR)', 'B58': 'Fee pass-through of ASP to revenue per unit (E5 D51)', 'B59': 'Revenue per unit growth — long-run anchor',
         'B60': 'Fade speed: share of the gap remaining after one year', 'B36': 'Disclosed revenue-per-unit and ASP growth (checks)', 'B14': 'check: service + vehicle sales − total revenue'}
for ref, text in short.items():
    old = r_[ref].value
    if old and old.strip() != text:
        r_[ref].value = text; note = r_[f'O{ref[1:]}'].value
        if ref not in ('B3', 'B14', 'B36', 'B47') and (not note or old.strip()[:30] not in note): r_[f'O{ref[1:]}'].value = (old.strip() + ('. ' + note if note else '')); r_[f'O{ref[1:]}'].font = font(C_NOTE, i=True, sz=10)
for row in (20, 25): group(r_, row, r_[f'B{row}'].value)
for row in (5, 9, 16, 32, 36):                                          # blue bars reach the notes column
    for c in range(2, 16): r_.cell(row, c).fill = fill(BLUE); r_.cell(row, c).border = BOT
for row in (42, 47, 53):
    for c in range(2, 16): r_.cell(row, c).fill = fill(GREY)
n = 0; skip_hist = {21, 22, 23}                                          # rows whose blank history cells feed the back-cast formulas
for row in range(6, r_.max_row + 1):
    b = r_[f'B{row}']
    if b.value is None or (b.fill is not None and b.fill.fill_type == 'solid') or row in (3, 4): continue
    cells = [r_.cell(row, c) for c in range(4, 14)]
    if not any(c.value is not None for c in cells): continue
    for c in cells:
        if c.value is None and not isinstance(c, MergedCell) and not (row in skip_hist and c.column <= 8):
            c.value = DASH; c.font = font(C_NOTE, i=True, sz=10); c.alignment = Alignment(horizontal='center'); n += 1
wb.save(P); print(f'saved {P.name}; RPM dashes {n}; freeze {r_.freeze_panes}')
