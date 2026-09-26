#!/usr/bin/env python3
"""Build model/CPRT_Intermediate.xlsx — the intermediate workbook the user copies from into "CPRT Model".

Every committed dataset in data/csv/ goes in as a Data_* tab with its provenance. The mechanism-A machinery
(fleet roll -> baseline TLF), the Solver calibration for R and P, the totaling-spread regression and the RPU chain
are LIVE FORMULAS off those data tabs. No named ranges (sheets survive "Move or Copy" into another workbook);
every cross-sheet reference is explicit.

PRESENTATION follows the Black Diamond (SOLS) model the user pointed at, plus the user's own quick tables:
  Garamond 11 · gridlines off · light-blue company bar / navy section bar / blue sub-section bars · navy year
  headers with A/E suffixes · thin vertical rules on data blocks · BLUE = hard-coded input · BLACK = formula ·
  GREEN = link to another sheet · PINK fill = calibrated parameter or key output · GREY fill = group header ·
  Cover tab with table of contents · coloured divider tabs between sections.

Usage:  ./.venv/bin/python scripts/build_intermediate_xlsx.py
Verify: ./.venv/bin/python scripts/verify_intermediate_xlsx.py   (pycel evaluation of the key outputs)
"""
import csv, pathlib, datetime, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

ROOT = pathlib.Path(__file__).resolve().parent.parent
CSV = ROOT / 'data/csv'; OUT = ROOT / 'model/CPRT_Intermediate.xlsx'
TODAY = datetime.date.today().isoformat()

# fitted parameters come from the header of data/csv/age_curves.csv (written by scripts/age_curves.py) so the workbook
# always matches the script
_hdr = " ".join(l for l in open(CSV / 'age_curves.csv', encoding='utf-8') if l.startswith('#'))
FIT = dict(
    k_cars=float(re.search(r"k_cars=([\d.]+)", _hdr).group(1)), k_lt=float(re.search(r"k_LT=([\d.]+)", _hdr).group(1)),
    r_slope=float(re.search(r"exp\(-([\d.]+)\*max\(a-6,0\)\)", _hdr).group(1)),
    pmin=float(re.search(r"P\(a\) = ([\d.]+) \+", _hdr).group(1)), pmax=float(re.search(r"P\(a\) = [\d.]+ \+ \(([\d.]+)-", _hdr).group(1)),
    pmid=float(re.search(r"exp\(-\(a-([\d.]+)\)/", _hdr).group(1)), pwid=float(re.search(r"exp\(-\(a-[\d.]+\)/([\d.]+)\)", _hdr).group(1)))
KMULT = 1.194
print("fitted parameters from age_curves.csv:", FIT)

# ---------------------------------------------------------------- palette & styles (SOLS conventions)
FONT = 'Garamond'
C_NAVY, C_BLUE, C_TITLE, C_LBLUE, C_PINK, C_GREY, C_YEL = '0E2841', '215E99', '0D5DB8', '95DCF7', 'F1CEEE', 'D0D0D0', 'FFFF00'
F_IN, F_FX, F_LINK, F_WHITE, F_NOTE = '0000FF', '000000', '008000', 'FFFFFF', '595959'
def font(color=F_FX, size=11, bold=False, italic=False, underline=None): return Font(name=FONT, size=size, bold=bold, italic=italic, color=color, underline=underline)
BLUE = font(F_IN); BLK = font(); GRN = font(F_LINK); BOLD = font(bold=True); NOTE = font(F_NOTE, 10, italic=True); WHITE_B = font(F_WHITE, bold=True)
def fill(c): return PatternFill('solid', fgColor=c)
NAVY, BLUEF, TITLEF, LBLUEF, PINK, GREY, YEL = map(fill, (C_NAVY, C_BLUE, C_TITLE, C_LBLUE, C_PINK, C_GREY, C_YEL))
THIN = Side(style='thin', color='808080')
PCT = '0.0%'; PCT2 = '0.00%'; DEC3 = '0.000'; DEC4 = '0.0000'; YR = '0'; YRA = '0"A"'; YRE = '0"E"'
ACTUAL_THRU = 2025

def put(ws, ref, v, f=BLK, fmt=None, fl=None, wrap=False, bold=False, align=None, valign=None):
    c = ws[ref]; c.value = v
    c.font = Font(name=f.name, size=f.size, bold=(bold or f.bold), italic=f.italic, color=f.color, underline=f.underline)
    if fmt: c.number_format = fmt
    if fl: c.fill = fl
    if wrap or align or valign: c.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical=valign or ('top' if wrap else None))
    return c

def widths(ws, spec):
    for col, w in spec.items(): ws.column_dimensions[col].width = w

def bar(ws, row, c0, c1, text, fl, f=WHITE_B, height=None, wrap=False):
    """a filled band across columns c0..c1 with text in the first cell"""
    for ci in range(c0, c1 + 1): ws.cell(row=row, column=ci).fill = fl
    put(ws, f'{L(c0)}{row}', text, f, fl=fl, wrap=wrap, valign='center')
    if c1 > c0: ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c1)
    if height: ws.row_dimensions[row].height = height

def skin(ws, title, subtitle, note, ncols, freeze=None, tab=None):
    """standard top block: company bar (row 1), navy tab title (row 2), italic note (row 3)"""
    ws.sheet_view.showGridLines = False
    bar(ws, 1, 1, ncols, 'Copart, Inc. (NASDAQ: CPRT) — intermediate build book', LBLUEF, font(C_NAVY, 11, bold=True), 18)
    bar(ws, 2, 1, max(1, ncols - 1) if subtitle else ncols, title, NAVY, font(F_WHITE, 12, bold=True), 20)
    if subtitle: put(ws, f'{L(ncols)}2', subtitle, font(F_WHITE, 10, italic=True), fl=NAVY, align='right', valign='center')
    put(ws, 'A3', note, NOTE, wrap=True); ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
    ws.row_dimensions[3].height = max(15, 13 * (1 + len(note) // (ncols * 11)))
    if freeze: ws.freeze_panes = freeze
    if tab: ws.sheet_properties.tabColor = tab

def header_row(ws, row, c0, c1, labels, fmt=None, wrap=True, height=None, center=True):
    for i, lab in enumerate(labels):
        put(ws, f'{L(c0+i)}{row}', lab, WHITE_B, fmt, fl=NAVY, wrap=wrap, align='center' if center else None, valign='center')
    for ci in range(c0 + len(labels), c1 + 1): ws.cell(row=row, column=ci).fill = NAVY
    if height: ws.row_dimensions[row].height = height

def year_header(ws, row, cols, years, link_from=None):
    for c, y in zip(cols, years):
        v = f'=Inputs!{c}{YROW}' if link_from == 'inputs' else y
        put(ws, f'{c}{row}', v, WHITE_B, (YRA if y <= ACTUAL_THRU else YRE), fl=NAVY, align='center')

def rules(ws, r0, r1, c0, c1, header_row_=None, total_rows=()):
    """thin vertical rules on a data block, bottom rule on its last row, box around the header"""
    for r in range(r0, r1 + 1):
        for ci in range(c0, c1 + 1):
            c = ws.cell(row=r, column=ci); b = c.border
            c.border = Border(left=THIN, right=THIN, top=(THIN if r in total_rows else b.top), bottom=(THIN if r == r1 else b.bottom))
    if header_row_:
        for ci in range(c0, c1 + 1):
            ws.cell(row=header_row_, column=ci).border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def read_csv(name):
    prov, rows = [], []
    with open(CSV / name, encoding='utf-8') as fh:
        for line in fh:
            if line.startswith('#'): prov.append(line[1:].strip())
            else: rows.append(line)
    rd = list(csv.reader(rows)); return prov, rd[0], rd[1:]

def num(v):
    if v is None or v == '': return None
    try:
        f = float(v); return int(f) if f.is_integer() and '.' not in v and 'e' not in v.lower() else f
    except ValueError: return v

def col_format(hdr, vals):
    nums = [v for v in vals if isinstance(v, (int, float))]
    if not nums: return None
    h = hdr.lower()
    if 'year' in h or h in ('fy', 'my') or all(isinstance(v, int) and 1900 <= v <= 2100 for v in nums): return YR
    if all(isinstance(v, int) for v in nums): return '#,##0'
    m = max(abs(v) for v in nums)
    return '#,##0.0' if m >= 100 else ('0.000' if m < 2 else '0.00')

# ---------------------------------------------------------------- data tab descriptions
DESC = {
 'reported_units.csv': "Copart quarterly YoY series, HAND-TRANSCRIBED from earnings-call transcripts (no filing discloses them). Tier 2. scripts/reported_series.py",
 'quarterly_pl.csv': "Copart quarterly P&L lines from 8-K Ex-99.1 press releases (SEC EDGAR). Tier 1.",
 'quarterly_margin.csv': "Copart quarterly gross margin and facility-operations % of revenue, from 8-K Ex-99.1. Tier 1.",
 'quarterly_segments.csv': "Copart US / international service and vehicle-sales revenue from 8-K, where disclosed. Tier 1.",
 'segment_service_rev_8k.csv': "US and international service revenue, 8-K Ex-99.1 (FY26). Tier 1.",
 'facility_ops_quarterly.csv': "Facility operations expense by quarter, 8-K Ex-99.1. Tier 1.",
 'sec_annual.csv': "Copart annual figures from 10-K XBRL (SEC EDGAR). Tier 1.",
 'elasticity_rebuild.csv': "RPU/ASP elasticity inputs, n=17 FY22Q4-FY26Q4: service & total revenue YoY (8-K), global units and global ASP YoY (transcripts, aligned). Implied RPU = (1+rev)/(1+units)-1. scripts/analysis_20260911.py",
 'units_decomp_panel_v2.csv': "Six-quarter US insurance units decomposition (Copart ex-CAT vs claims x TLF pool) under three claims denominators. scripts/units_decomp_panel.py; findings.md Addendum 15/15a",
 'decomposition.csv': "Matched-denominator decomposition, CY2025 and FQ4 FY26 windows. scripts/analysis_20260911.py",
 'tlf_calibration.csv': "Regression outputs: dTLF(pp, YoY) = a + b x spread(t-1), CCC quarterly 2019Q1-2025Q3 (n=27). scripts/analysis_20260911.py",
 'totaling_spread_quarterly.csv': "Totaling spread = repair CPI YoY - used-car CPI YoY (BLS CUUR0000SETD minus CUUR0000SETA02, NSA), quarterly average of monthly values. months<3 = partial quarter. scripts/analysis_20260911.py",
 'cprt_cpi_three_series.csv': "BLS CPI flat files (download.bls.gov/pub/time.series/cu): SETA02 used cars & trucks, SETD motor vehicle maintenance & repair, SETE motor vehicle insurance; SA (CUSR) and NSA (CUUR); footnote code X marks the 2025 appropriations-lapse gap (left blank, never interpolated).",
 'fasttrack_collision_claims_cw.csv': "ISS Fast Track quarterly collision-claim-count headlines as reported by CollisionWeek (collisionweek.com/tag/losses/, headline + teaser only). Headline figures carry qualifiers.",
 'pgr_monthly_pif.csv': "Progressive personal-auto policies in force, monthly 8-K EX-99 (CIK 0000080661), SEC EDGAR.",
 'geico_frequency_series.csv': "Berkshire Hathaway 10-Q/10-K sentences on GEICO claims frequency, SEC EDGAR (CIK 0001067983).",
 'rba_automotive_series.csv': "RB Global (IAA parent) automotive lots / GTV / take rate from 8-K Ex-99.1 (CIK 0001046102). 5 of 18 quarters are pro forma - see basis column.",
 'duopoly_compare.csv': "RB Global automotive lots YoY vs Copart US insurance units YoY by matched quarter.",
 'duopoly_daily.csv': "Daily listed US inventory: Copart (lot.xml sitemap union) vs IAA (vehicledetail sitemap). scripts/job1_snapshot.py, job1_iaa.py. Listed inventory, not yard inventory.",
 'yard_panel_us.csv': "Copart US yards and dated sale events per sitemap snapshot (Wayback 2022-2025 + live).",
 'cadence_fixed.csv': "Sale-event cadence per yard per snapshot.",
 'lots_per_sale_event.csv': "Listed lots (union, overlap-gated) per 7-day sale events; events per yard per week.",
 'backtest_inventory_v2.csv': "Sitemap inventory YoY vs Copart reported US inventory YoY, overlap<=1% gate. scripts/backtest_inventory_v2.py",
 'fred_TOTALNSA.csv': "FRED TOTALNSA: light weight vehicle sales, autos and light trucks, thousands of units, monthly NSA (BEA).",
 'fred_LTRUCKNSA.csv': "FRED LTRUCKNSA: light weight vehicle sales, light trucks, thousands of units, monthly NSA (BEA). Fetched 2026-09-14.",
 'fred_HTRUCKSNSA.csv': "FRED HTRUCKSNSA: motor vehicle retail sales, heavy weight trucks (>14,000 lb GVW), thousands of units, monthly NSA (BEA). Fetched 2026-09-15. Ward's 'heavy' is >10,000 lb, so the 2021 overlap ratio (~1.8) is a class-3 definitional gap.",
 'capex_decomp.csv': "Copart capex decomposition from 10-K PP&E roll-forward (XBRL).",
 'owner_earnings.csv': "Owner-earnings bridge (EBIT less non-land capex above D&A), 10-K XBRL.",
 'ornl_tedb40_miles_by_age.csv': "EPA annual vehicle-miles-of-travel by age (EPA-420-D-16-900, July 2016), republished as ORNL TEDB Ed.40 Table 3.14.",
 'age_curves.csv': "Per-age output of scripts/age_curves.py: EPA survival, stretched survival, miles, R, P, 2024 fleet/claims/TL shares.",
 'age_curves_validation.csv': "Every fit target, out-of-sample check, survival check and sensitivity row from scripts/age_curves.py.",
}
DATA_TABS = [  # (sheet, csv, display title)
 ('Data_ReportedUnits', 'reported_units.csv', 'Copart reported unit, inventory and ASP growth by fiscal quarter (transcripts)'),
 ('Data_QuarterlyPL', 'quarterly_pl.csv', 'Copart quarterly P&L lines (8-K)'), ('Data_QuarterlyMargin', 'quarterly_margin.csv', 'Copart quarterly gross margin and facility-ops ratio (8-K)'),
 ('Data_Segments', 'quarterly_segments.csv', 'Copart US / international revenue split (8-K)'), ('Data_SegmentSvc8K', 'segment_service_rev_8k.csv', 'US and international service revenue (8-K, FY26)'),
 ('Data_FacilityOps', 'facility_ops_quarterly.csv', 'Facility operations expense by quarter (8-K)'), ('Data_SECAnnual', 'sec_annual.csv', 'Copart annual figures (10-K XBRL)'),
 ('Data_Elasticity', 'elasticity_rebuild.csv', 'RPU vs ASP elasticity inputs, n = 17 quarters'), ('Data_CCC_TLF_Annual', 'ccc_tlf_annual.csv', 'CCC total-loss frequency, annual 2013–2025'),
 ('Data_CCC_TLF_Quarterly', 'ccc_tlf_quarterly.csv', 'CCC total-loss frequency, quarterly 2018Q1–2025Q3'), ('Data_CPI', 'cprt_cpi_three_series.csv', 'BLS CPI: used cars, vehicle repair, vehicle insurance (SA and NSA), monthly'),
 ('Data_Spread', 'totaling_spread_quarterly.csv', 'Totaling spread by quarter (repair CPI YoY − used-car CPI YoY)'), ('Data_TLFCalibration', 'tlf_calibration.csv', 'Spread-regression outputs (script)'),
 ('Data_UnitsPanel', 'units_decomp_panel_v2.csv', 'Six-quarter units decomposition panel'), ('Data_Decomposition', 'decomposition.csv', 'Matched-denominator decomposition windows'),
 ('Data_FastTrack', 'fasttrack_collision_claims_cw.csv', 'ISS Fast Track collision-claim headlines (CollisionWeek)'), ('Data_PGR_PIF', 'pgr_monthly_pif.csv', 'Progressive personal-auto policies in force, monthly'),
 ('Data_GEICO', 'geico_frequency_series.csv', 'GEICO claims-frequency statements (Berkshire filings)'), ('Data_RBA', 'rba_automotive_series.csv', 'RB Global automotive lots, GTV, take rate'),
 ('Data_DuopolyCompare', 'duopoly_compare.csv', 'RB Global vs Copart unit growth by quarter'), ('Data_DuopolyDaily', 'duopoly_daily.csv', 'Copart vs IAA listed US inventory, daily'),
 ('Data_Yards', 'yard_panel_us.csv', 'Copart US yards and sale events per snapshot'), ('Data_Cadence', 'cadence_fixed.csv', 'Sale-event cadence per yard'),
 ('Data_LotsPerEvent', 'lots_per_sale_event.csv', 'Listed lots per sale event'), ('Data_BacktestInv', 'backtest_inventory_v2.csv', 'Sitemap inventory vs reported inventory back-test'),
 ('Data_FRED_TOTALNSA', 'fred_TOTALNSA.csv', 'FRED TOTALNSA — light-vehicle sales, monthly'), ('Data_FRED_LTRUCKNSA', 'fred_LTRUCKNSA.csv', 'FRED LTRUCKNSA — light-truck sales, monthly'),
 ('Data_FRED_HTRUCKSNSA', 'fred_HTRUCKSNSA.csv', 'FRED HTRUCKSNSA — heavy-truck sales, monthly'), ('Data_EPA_Miles', 'ornl_tedb40_miles_by_age.csv', 'Miles(age): annual miles per vehicle by age (EPA)'),
 ('Data_AgeCurvesPy', 'age_curves.csv', 'Age curves as computed by the script (reference copy)'), ('Data_AgeCurvesValid', 'age_curves_validation.csv', 'Age-curve validation table (script)'),
 ('Data_Capex', 'capex_decomp.csv', 'Copart capex decomposition (10-K)'), ('Data_OwnerEarnings', 'owner_earnings.csv', 'Owner-earnings bridge (10-K)'),
]

def add_data_sheet(wb, name, fname, title):
    prov, hdr, rows = read_csv(fname); ws = wb.create_sheet(name); nc = max(6, len(hdr))
    skin(ws, title, f'data/csv/{fname} · {len(rows)} rows', (" | ".join(prov) if prov else DESC.get(fname, "see PROVENANCE.md")) + "   ·   BLUE = raw data as committed; edit the CSV and rebuild rather than typing here.", nc, freeze='B5')
    header_row(ws, 4, 1, nc, hdr, height=30)
    cols = list(zip(*[[num(v) for v in r] for r in rows])) if rows else [[] for _ in hdr]
    fmts = [col_format(h, list(cv)) for h, cv in zip(hdr, cols)]
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            put(ws, f'{L(j+1)}{5+i}', num(v), BLUE, fmts[j])
    rules(ws, 5, 4 + len(rows), 1, len(hdr), header_row_=4)
    for j, h in enumerate(hdr):
        w = max([len(h) // 2 + 4] + [min(40, len(str(r[j])) + 2) for r in rows[:200]]) if rows else len(h) + 2
        ws.column_dimensions[L(j+1)].width = min(44, max(9, w))
    return ws

# ---------------------------------------------------------------- workbook
wb = Workbook(); wcv = wb.active; wcv.title = 'Cover'
Y0, Y1 = 2013, 2030; YEARS = list(range(Y0, Y1 + 1)); NY = len(YEARS); YC = [L(3 + i) for i in range(NY)]   # C..T
AGES = list(range(0, 46)); NA = len(AGES)
MY0, MY1 = 1970, 2030; MYS = list(range(MY0, MY1 + 1))

# ================================================================ Inputs
wi = wb.create_sheet('Inputs')
skin(wi, 'Inputs — every parameter of the fleet roll, once', 'the only tab the formula sheets read parameters from',
     'BLUE = hard-coded input (change freely) · PINK fill = calibrated parameter (fitted on the Calibration tab / scripts/age_curves.py — paste new Solver results here) · every other tab links here in GREEN.', 4, freeze='A5')
IN = {}
def in_section(row, text): bar(wi, row, 1, 4, text, BLUEF, WHITE_B, 16)
def in_row(row, key, label, value, fmt, units, note, calibrated=False):
    put(wi, f'A{row}', label, BLK); put(wi, f'B{row}', value, BLUE, fmt, fl=(PINK if calibrated else None), align='right'); put(wi, f'C{row}', units, NOTE, align='center'); put(wi, f'D{row}', note, NOTE, wrap=True)
    IN[key] = f'Inputs!$B${row}'
_, _, sales_rows = read_csv('light_vehicle_sales_by_year.csv'); S25 = [r for r in sales_rows if r[0] == '2025'][0]
header_row(wi, 4, 1, 4, ['Parameter', 'Value', 'Units', 'Source / note'], wrap=False, center=False)
in_section(5, 'Survival S(age) — EPA schedule stretched along the age axis')
in_row(6,  'k_cars', 'Stretch k, cars, 2013',           round(FIT['k_cars'] / KMULT, 3), DEC3, 'x', "fitted to the IHS/Polk 2013 census of cars by single year of age (ORNL TEDB40 T3.11); S_k(a) = S_EPA(a/k). docs/AGE_CURVES.md §2.2", True)
in_row(7,  'k_lt',   'Stretch k, light trucks, 2013',   round(FIT['k_lt'] / KMULT, 3), DEC3, 'x', "fitted to the IHS 2013 census of trucks (T3.12, includes heavy). §2.2", True)
in_row(8,  'k_mult', 'k multiplier by drift-end year',   KMULT, DEC3, 'x', "fitted so the 2024 roll hits the VIO anchor and 66% aged 7+. 289M anchor is unverified (from memory); Experian 292.1M (on disk) gives 1.204 — same downstream. §2.4", True)
in_row(9,  'drift0', 'Drift start year',                 2013, YR, 'year', "k = 2013 value up to here")
in_row(10, 'drift1', 'Drift end year',                   2024, YR, 'year', "k reaches k_2013 × multiplier here; FLAT after (base case). Sensitivity: set 2030 to keep survival improving")
in_section(12, 'Claim frequency R(age) — claims per vehicle on the road, relative to a 1–6-year-old. FITTED, not measured: the slope below is the one free parameter')
in_row(13, 'expo0', 'Exposure factor, age 0',            0.5, '0.00', 'x', "a model-year-t vehicle is on the road ~half of calendar year t. Necessary: without it the ≤3-yr repairable share fits at 36% vs CCC ~30%. §3.2")
in_row(14, 'kink',  'Kink age',                          6, YR, 'yrs', "R is flat through this age (a free 0–6 slope is not identified by the CCC targets). §3.3")
in_row(15, 'slope', 'Decline per year of age after the kink', FIT['r_slope'], DEC4, 'per yr', "R(a) = e(a)·exp(−slope·max(a−kink,0)). Solver-fitted with P to eight CCC 2024 statistics. §3.3", True)
in_section(17, 'Total-loss propensity P(age) — share of claims declared a total loss, logistic in age')
in_row(18, 'pmin',  'Floor (pmin)',                      FIT['pmin'], DEC3, 'share', "P(a) = pmin + (pmax−pmin)/(1+exp(−(a−mid)/width)). Solver-fitted. §4.3", True)
in_row(19, 'pmax',  'Ceiling (pmax)',                    FIT['pmax'], DEC3, 'share', "", True)
in_row(20, 'pmid',  'Midpoint age (mid)',                FIT['pmid'], '0.00', 'yrs', "", True)
in_row(21, 'pwid',  'Width',                             FIT['pwid'], '0.00', 'yrs', "", True)
in_section(23, 'Cohort sizes — new-vehicle sales assumed for model years 2026+')
in_row(24, 'fut_cars', 'Cars',                           round(float(S25[1]), 1), '#,##0', 'thousands / yr', "ASSUMED = 2025 level (Data_Sales). Flex for a sales-recovery case")
in_row(25, 'fut_lt',   'Light trucks',                   round(float(S25[2]), 1), '#,##0', 'thousands / yr', "ASSUMED = 2025 level")
in_section(27, 'Fleet-count anchors used to fit the survival drift (checked on FleetByAge)')
in_row(28, 'vio',   'Light vehicles in operation, end-2024', 289, '#,##0', 'millions', "S&P Global Mobility 2025 avg-age release — NOT ON DISK, confirm; Experian Q3-2024 = 292.1M (Crash Course 2025/Q1)")
in_row(29, 'sh7',   'Share of fleet aged 7+, 2024',      0.66, PCT, 'share', "'66% of vehicles in operation are seven years or older' — S&P via CCC Crash Course 2024/Q4")
YROW, KMROW, KCROW, KLROW = 33, 34, 35, 36
bar(wi, 31, 1, 2 + NY, 'Survival stretch by calendar year (formulas) — these year headers drive every year-indexed tab', BLUEF, WHITE_B, 16)
put(wi, f'A{YROW}', 'Calendar year', WHITE_B, fl=NAVY); put(wi, f'B{YROW}', 'units', WHITE_B, fl=NAVY, align='center')
put(wi, f'A{KMROW}', 'k multiplier', BLK); put(wi, f'A{KCROW}', 'k cars', BLK); put(wi, f'A{KLROW}', 'k light trucks', BLK)
for r in (KMROW, KCROW, KLROW): put(wi, f'B{r}', 'x', NOTE, align='center')
for i, y in enumerate(YEARS):
    c = YC[i]
    put(wi, f'{c}{YROW}', y, WHITE_B, (YRA if y <= ACTUAL_THRU else YRE), fl=NAVY, align='center')
    put(wi, f'{c}{KMROW}', f'=1+({IN["k_mult"]}-1)*MIN(MAX({c}{YROW}-{IN["drift0"]},0),{IN["drift1"]}-{IN["drift0"]})/({IN["drift1"]}-{IN["drift0"]})', BLK, DEC4)
    put(wi, f'{c}{KCROW}', f'={IN["k_cars"]}*{c}{KMROW}', BLK, DEC4); put(wi, f'{c}{KLROW}', f'={IN["k_lt"]}*{c}{KMROW}', BLK, DEC4)
rules(wi, KMROW, KLROW, 3, 2 + NY, header_row_=YROW)
put(wi, 'A38', 'A = fleet built from actual sales; E = 2026+ model years use the assumed sales above. Sensitivity: set B10 = 2030 to let survival keep improving; set B8 = 1.0 for the raw EPA schedule. Baseline TLF drift moves 0.04–0.18pp/yr across that range; P buckets do not move (docs/AGE_CURVES.md §6).', NOTE, wrap=True)
wi.merge_cells('A38:H38'); wi.row_dimensions[38].height = 42
for r0, r1 in ((6, 10), (13, 15), (18, 21), (24, 25), (28, 29)): rules(wi, r0, r1, 1, 4)
widths(wi, {'A': 42, 'B': 12, 'C': 13, 'D': 96})
for c in YC: wi.column_dimensions[c].width = 8

# ================================================================ Data_Sales (specific layout: MY 1970..2030)
wsl = wb.create_sheet('Data_Sales')
prov, hdr, rows = read_csv('light_vehicle_sales_by_year.csv')
skin(wsl, 'Sales(MY): new vehicle sales by model year — cohort sizes for the fleet roll', 'thousands of vehicles',
     "Unit: thousands (8,321 = 8,321,000 cars). Ward's reports to the single vehicle, so the source carries decimals (kept in the cells, hidden by the format). Ward's via ORNL TEDB Ed.40 Table 3.6 for 1970–2021; FRED TOTALNSA / LTRUCKNSA / HTRUCKSNSA rescaled to Ward's basis on the 2021 overlap for 2022–2025 (estimates); 2026+ = Inputs (green). Heavy trucks (>10,000 lb) are NOT in the fleet roll — the column exists for the 2013 census check only. data/csv/light_vehicle_sales_by_year.csv", 5, freeze='A5')
header_row(wsl, 4, 1, 5, ['Model year', 'Cars (thousands)', 'Light trucks (thousands)', 'Heavy trucks (thousands)', 'Source'], height=30)
byy = {int(r[0]): r for r in rows}; SR0 = 5
for i, my in enumerate(MYS):
    r = SR0 + i; put(wsl, f'A{r}', my, BLUE, YR, align='center')
    if my in byy:
        put(wsl, f'B{r}', round(float(byy[my][1]), 1), BLUE, '#,##0'); put(wsl, f'C{r}', round(float(byy[my][2]), 1), BLUE, '#,##0')
        put(wsl, f'D{r}', round(float(byy[my][3]), 1), BLUE, '#,##0'); put(wsl, f'E{r}', byy[my][4], NOTE)
    else:
        put(wsl, f'B{r}', f'={IN["fut_cars"]}', GRN, '#,##0'); put(wsl, f'C{r}', f'={IN["fut_lt"]}', GRN, '#,##0'); put(wsl, f'D{r}', None, BLUE); put(wsl, f'E{r}', 'ASSUMED — Inputs (2025 level); heavy trucks not modelled', NOTE)
SR1 = SR0 + len(MYS) - 1
rules(wsl, SR0, SR1, 1, 5, header_row_=4)
widths(wsl, {'A': 12, 'B': 16, 'C': 20, 'D': 20, 'E': 60})
SALES_Y = f"Data_Sales!$A${SR0}:$A${SR1}"; SALES_C = f"Data_Sales!$B${SR0}:$B${SR1}"; SALES_L = f"Data_Sales!$C${SR0}:$C${SR1}"

# ================================================================ Data_EPA_Survival
wse = wb.create_sheet('Data_EPA_Survival')
prov, hdr, rows = read_csv('ornl_tedb40_survival_by_age.csv')
skin(wse, 'S(age): survival rates for cars and light trucks by vehicle age (EPA, 2016)', 'share of a model-year cohort still registered at each age',
     'Primary: U.S. EPA, Draft Technical Assessment Report, Midterm Evaluation of Light-Duty GHG / CAFE Standards MY2022–2025, EPA-420-D-16-900, July 2016 — republished as ORNL Transportation Energy Data Book Ed.40 Table 3.15. Values verbatim. The Survival tab stretches this schedule along the age axis by k(t) from Inputs.', 4, freeze='A5')
header_row(wse, 4, 1, 4, ['Vehicle age (years)', 'Survival rate — cars', 'Survival rate — light trucks', 'Source'], height=30)
for i, r in enumerate(rows):
    put(wse, f'A{5+i}', int(r[0]), BLUE, YR, align='center'); put(wse, f'B{5+i}', float(r[1]), BLUE, DEC3, align='center'); put(wse, f'C{5+i}', float(r[2]), BLUE, DEC3, align='center')
put(wse, 'D5', 'EPA schedule (ORNL TEDB Ed.40 T3.15), 2016', NOTE)
rules(wse, 5, 4 + len(rows), 1, 4, header_row_=4)
EPA_C = "Data_EPA_Survival!$B$5:$B$36"; EPA_L = "Data_EPA_Survival!$C$5:$C$36"
widths(wse, {'A': 18, 'B': 20, 'C': 24, 'D': 44})

# ================================================================ Data_CCC_Targets
wt = wb.create_sheet('Data_CCC_Targets')
skin(wt, 'CCC targets — the statistics R(age) and P(age) are fitted to (FIT) and tested against (OOS), and the fleet-count anchors (S)', 'quotes verbatim from the ungated report pages',
     'Tolerance = the error that costs one unit of Solver loss on the Calibration tab. Pages fetched 2026-09-11 (PROVENANCE.md §5). FIT rows must stay at the top in this order — Calibration references them by row.', 8, freeze='A5')
header_row(wt, 4, 1, 8, ['Use', 'Year', 'Key', 'Statistic', 'Value', 'Tolerance', 'Quote', 'Page'], height=30)
TGT = [
 ('FIT', 2024, 'tlf',    'Total-loss frequency, all loss categories', 0.223, 0.005, 'printed data label, chart "Total Loss Frequency Remains High as Vehicle Values Trend Up" (2024 = 22.3%)', 'https://www.cccis.com/reports/crash-course-2026'),
 ('FIT', 2024, 'tl7',    'Share of total-loss valuations on vehicles 7+ yrs', 0.72, 0.01, 'almost 72% of valuations across all loss categories are for vehicles 7 years or older', 'https://www.cccis.com/reports/crash-course-2024/q4'),
 ('FIT', 2024, 'rp7',    'Share of repairable claims on vehicles 7+ yrs', 0.45, 0.01, 'vehicles seven years or older now make up nearly 45% of all repairable claims, up from 35% in 2019', 'https://www.cccis.com/reports/crash-course-2024/q4'),
 ('FIT', 2024, 'rp3',    'Share of repairable claims on vehicles ≤3 yrs (all fuel)', 0.30, 0.02, 'Only 26.3% of repairable ICE vehicles are three years or newer (EVs 79.4% / hybrids 60.3% ≤3 yrs → all-fuel ≈30%)', 'https://www.cccis.com/reports/crash-course-2024/q4'),
 ('FIT', 2024, 'age_cl', 'Average age of vehicles in claims (yrs)', 7.6, 0.15, 'For claims, the average age of vehicles has increased to 7.6 years – up from 6.9 years in 2020', 'https://www.cccis.com/reports/crash-course-2025/q1'),
 ('FIT', 2024, 'age_rp', 'Average age of repairable vehicles (yrs)', 6.8, 0.15, 'The average age of repairable vehicles was 6.8 years in 2024 (up from 6.1 years old in 2020)', 'https://www.cccis.com/reports/crash-course-2025/q1'),
 ('FIT', 2024, 'age_tl', 'Average age of total-loss vehicles (yrs)', 10.6, 0.15, '... and 10.6 years for total loss vehicles (up from 10.0 years old in 2020)', 'https://www.cccis.com/reports/crash-course-2025/q1'),
 ('FIT', 2024, 'p03',    'Total-loss rate for claims on vehicles ≤3 yrs', 0.10, 0.01, 'For claims that are three years old or newer, 1 in 10 are flagged as a total loss by the insurer', 'https://www.cccis.com/reports/crash-course-2026'),
 ('OOS', 2019, 'rp7',    'Share of repairable claims on vehicles 7+ yrs', 0.35, None, 'up from 35% in 2019', 'https://www.cccis.com/reports/crash-course-2024/q4'),
 ('OOS', 2019, 'rp3',    'Share of repairable claims on vehicles ≤3 yrs', 0.385, None, 'vehicles 3 years or newer represented 38.5% of the repairable mix in 2019', 'https://www.cccis.com/reports/crash-course-2025/q1'),
 ('OOS', 2020, 'age_cl', 'Average age of vehicles in claims (yrs)', 6.9, None, 'up from 6.9 years in 2020', 'https://www.cccis.com/reports/crash-course-2025/q1'),
 ('OOS', 2020, 'age_rp', 'Average age of repairable vehicles (yrs)', 6.1, None, 'up from 6.1 years old in 2020', 'https://www.cccis.com/reports/crash-course-2025/q1'),
 ('OOS', 2020, 'age_tl', 'Average age of total-loss vehicles (yrs)', 10.0, None, 'up from 10.0 years old in 2020', 'https://www.cccis.com/reports/crash-course-2025/q1'),
 ('OOS', 2025, 'tl7',    'Share of total-loss valuations on vehicles 7+ yrs', 0.72, None, 'over 72% of total loss valuations are on vehicles 7 years or older', 'https://www.cccis.com/reports/crash-course-2025/q4'),
 ('OOS', 2025, 'rp7',    'Share of repairable claims on vehicles 7+ yrs', 0.46, None, 'almost 46% of repairable vehicles are 7 years or older', 'https://www.cccis.com/reports/crash-course-2025/q4'),
 ('OOS', 2019, 'tlf',    'TLF all loss categories (demographics-only model vs actual)', 0.192, None, 'CCC annual 2019 = 19.2%; the gap is the non-demographic LEVEL (spread cycle, technology, filing)', 'Data_CCC_TLF_Annual'),
 ('OOS', 2020, 'tlf',    'TLF all loss categories', 0.206, None, 'CCC annual 2020 = 20.6%', 'Data_CCC_TLF_Annual'),
 ('OOS', 2025, 'tlf',    'TLF all loss categories', 0.231, None, 'CCC annual 2025 = 23.1%', 'Data_CCC_TLF_Annual'),
 ('S',   2024, 'vio',    'Light vehicles in operation (M)', 289, None, 'S&P 2025 release — NOT on disk (from memory); Experian 292.1M at Q3-2024 is on disk (Crash Course 2025/Q1)', 'https://www.cccis.com/reports/crash-course-2025/q1'),
 ('S',   2024, 'sh7',    'Share of fleet aged 7+', 0.66, None, '66% of vehicles in operation are seven years or older (S&P Mobility)', 'https://www.cccis.com/reports/crash-course-2024/q4'),
 ('S',   2025, 'viogr',  'VIO growth 2020→Q3-2025', 0.051, None, 'Vehicles in operation reached almost 296 million as of Q3 2025 (Experian), +5.1% relative to 2020', 'https://www.cccis.com/reports/crash-course-2026'),
]
for i, t in enumerate(TGT):
    r = 5 + i
    for j, v in enumerate(t):
        fmt = None
        if j == 4: fmt = ('0.0' if t[2].startswith('age') else ('#,##0' if t[2] == 'vio' else PCT))
        if j == 5: fmt = ('0.00' if t[2].startswith('age') else PCT2)
        put(wt, f'{L(j+1)}{r}', v, BLUE if j in (4, 5) else BLK, fmt, wrap=(j == 6), align=('center' if j in (0, 1, 2, 4, 5) else None))
    if t[0] == 'FIT': wt[f'A{r}'].fill = PINK
TR0, TR1 = 5, 4 + len(TGT)
rules(wt, TR0, TR1, 1, 8, header_row_=4)
widths(wt, {'A': 6, 'B': 7, 'C': 8, 'D': 48, 'E': 9, 'F': 10, 'G': 84, 'H': 48})

# ================================================================ Survival: S(a, t)
wsv = wb.create_sheet('Survival')
CR0, CR1 = 6, 6 + NA - 1; LR0, LR1 = CR1 + 4, CR1 + 4 + NA - 1
skin(wsv, 'Survival by age and calendar year — EPA schedule stretched by k(t)', 'share of each model-year cohort still on the road',
     'S(a,t) = S_EPA(a / k(t)), linearly interpolated between the table’s integer ages, zero beyond age 31·k. k(t) links from Inputs (green). Ages 0–45; S is ~0 past 40. Cars block rows 6–51, light trucks rows 55–100.', 2 + NY, freeze=f'C{CR0}')
# row 3 doubles as the year header on this tab (formulas reference Survival!C3:T3)
wsv.unmerge_cells(start_row=3, start_column=1, end_row=3, end_column=2 + NY)
put(wsv, 'A3', 'Calendar year →', WHITE_B, fl=NAVY); put(wsv, 'B3', 'units', WHITE_B, fl=NAVY, align='center'); wsv.row_dimensions[3].height = 15
year_header(wsv, 3, YC, YEARS, link_from='inputs')
put(wsv, 'A4', 'k cars (Inputs)', BLK); put(wsv, 'B4', 'x', NOTE, align='center')
put(wsv, f'A{CR0-1}', 'CARS — age ↓', BOLD, fl=GREY); put(wsv, f'B{CR0-1}', 'years', NOTE, fl=GREY, align='center')
put(wsv, f'A{LR0-2}', 'k light trucks (Inputs)', BLK); put(wsv, f'B{LR0-2}', 'x', NOTE, align='center')
put(wsv, f'A{LR0-1}', 'LIGHT TRUCKS — age ↓', BOLD, fl=GREY); put(wsv, f'B{LR0-1}', 'years', NOTE, fl=GREY, align='center')
for i, y in enumerate(YEARS):
    c = YC[i]
    put(wsv, f'{c}4', f'=Inputs!{c}{KCROW}', GRN, DEC4); put(wsv, f'{c}{LR0-2}', f'=Inputs!{c}{KLROW}', GRN, DEC4)
    put(wsv, f'{c}{CR0-1}', f'={c}3', BOLD, YR, fl=GREY, align='center'); put(wsv, f'{c}{LR0-1}', f'={c}3', BOLD, YR, fl=GREY, align='center')
for i, a in enumerate(AGES):
    put(wsv, f'A{CR0+i}', a, BLUE, YR, align='center'); put(wsv, f'A{LR0+i}', a, BLUE, YR, align='center')
    for c in YC:
        x = f'($A{CR0+i}/{c}$4)'
        put(wsv, f'{c}{CR0+i}', f'=IF({x}>=31,0,INDEX({EPA_C},INT({x})+1)+({x}-INT({x}))*(INDEX({EPA_C},INT({x})+2)-INDEX({EPA_C},INT({x})+1)))', BLK, DEC3)
        x = f'($A{LR0+i}/{c}${LR0-2})'
        put(wsv, f'{c}{LR0+i}', f'=IF({x}>=31,0,INDEX({EPA_L},INT({x})+1)+({x}-INT({x}))*(INDEX({EPA_L},INT({x})+2)-INDEX({EPA_L},INT({x})+1)))', BLK, DEC3)
rules(wsv, CR0, CR1, 3, 2 + NY, header_row_=CR0 - 1); rules(wsv, LR0, LR1, 3, 2 + NY, header_row_=LR0 - 1)
widths(wsv, {'A': 24, 'B': 7})
for c in YC: wsv.column_dimensions[c].width = 8
SV_YEARS = f"Survival!$C$3:${YC[-1]}$3"; SV_CARS = f"Survival!$C${CR0}:${YC[-1]}${CR1}"; SV_LT = f"Survival!$C${LR0}:${YC[-1]}${LR1}"

# ================================================================ Fleet: model year x calendar year
wf = wb.create_sheet('Fleet')
FR0 = 6; FR1 = FR0 + len(MYS) - 1
skin(wf, 'Fleet roll — vehicles on the road by model year and calendar year', 'thousands of vehicles, cars + light trucks',
     'cell = sales_cars(MY) × S_cars(t−MY, t) + sales_LT(MY) × S_LT(t−MY, t). Blank where the model year is after the calendar year. Bottom row = total vehicles in operation. A/E on the year headers: fleet from actual sales / from assumed 2026+ sales.', 2 + NY, freeze=f'C{FR0}')
put(wf, 'A4', 'Model year ↓  /  calendar year →', WHITE_B, fl=NAVY); put(wf, 'B4', 'units', WHITE_B, fl=NAVY, align='center')
year_header(wf, 4, YC, YEARS, link_from='inputs')
put(wf, 'A5', '', fl=GREY); put(wf, 'B5', 'thousands', NOTE, fl=GREY, align='center')
for c in YC: wf[f'{c}5'].fill = GREY
for i, my in enumerate(MYS):
    r = FR0 + i; put(wf, f'A{r}', my, BLUE, YR, align='center')
    for c in YC:
        age = f'({c}$4-$A{r})'
        put(wf, f'{c}{r}', f'=IF({c}$4<$A{r},"",IF({age}>45,0,INDEX({SALES_C},MATCH($A{r},{SALES_Y},0))*INDEX({SV_CARS},{age}+1,MATCH({c}$4,{SV_YEARS},0))+INDEX({SALES_L},MATCH($A{r},{SALES_Y},0))*INDEX({SV_LT},{age}+1,MATCH({c}$4,{SV_YEARS},0))))', BLK, '#,##0')
TOTR = FR1 + 2
put(wf, f'A{TOTR}', 'Total vehicles in operation', BOLD, fl=PINK); put(wf, f'B{TOTR}', 'thousands', NOTE, fl=PINK, align='center')
for c in YC: put(wf, f'{c}{TOTR}', f'=SUM({c}{FR0}:{c}{FR1})', BOLD, '#,##0', fl=PINK)
rules(wf, FR0, FR1, 3, 2 + NY, header_row_=4); rules(wf, TOTR, TOTR, 3, 2 + NY)
widths(wf, {'A': 30, 'B': 10})
for c in YC: wf.column_dimensions[c].width = 9
FL_YEARS = f"Fleet!$C$4:${YC[-1]}$4"; FL_MY = f"Fleet!$A${FR0}:$A${FR1}"; FL_MAT = f"Fleet!$C${FR0}:${YC[-1]}${FR1}"

# ================================================================ FleetByAge
wa = wb.create_sheet('FleetByAge')
AR0 = 6; AR1 = AR0 + NA - 1
skin(wa, 'Fleet by age and calendar year — and the fleet statistics the survival drift is checked against', 'thousands of vehicles',
     'fleet(a,t) = Fleet[MY = t−a, t]. Statistics below the block: vehicles in operation, share aged 7+ and 15+, count aged ≤6, average age. Compare with the S rows on Data_CCC_Targets. Average age is NOT comparable to S&P’s 12.8 — S&P’s figure needs a 15+ tail averaging ~27 yrs (docs/AGE_CURVES.md §2.3); calibrate to counts.', 2 + NY, freeze=f'C{AR0}')
put(wa, 'A4', 'Age ↓  /  calendar year →', WHITE_B, fl=NAVY); put(wa, 'B4', 'age + 0.5', WHITE_B, fl=NAVY, align='center')
year_header(wa, 4, YC, YEARS, link_from='inputs')
put(wa, 'A5', '', fl=GREY); put(wa, 'B5', 'helper', NOTE, fl=GREY, align='center')
for c in YC: wa[f'{c}5'].fill = GREY
for i, a in enumerate(AGES):
    r = AR0 + i; put(wa, f'A{r}', a, BLUE, YR, align='center'); put(wa, f'B{r}', f'=A{r}+0.5', BLK, '0.0', align='center')
    for c in YC:
        put(wa, f'{c}{r}', f'=IF({c}$4-$A{r}<{MY0},0,INDEX({FL_MAT},MATCH({c}$4-$A{r},{FL_MY},0),MATCH({c}$4,{FL_YEARS},0)))', BLK, '#,##0')
ST = {'vio': AR1 + 2, 'sh7': AR1 + 3, 'sh15': AR1 + 4, 'le6': AR1 + 5, 'avg': AR1 + 6}
lbl = {'vio': ('Vehicles in operation', 'millions'), 'sh7': ('Share aged 7+', '%'), 'sh15': ('Share aged 15+', '%'), 'le6': ('Count aged ≤6', 'millions'), 'avg': ('Average age (t − MY + 0.5)', 'years')}
for k, r in ST.items():
    key = k in ('vio', 'sh7')
    put(wa, f'A{r}', lbl[k][0], BOLD if key else BLK, fl=(PINK if key else None)); put(wa, f'B{r}', lbl[k][1], NOTE, align='center', fl=(PINK if key else None))
for c in YC:
    rng = f'{c}{AR0}:{c}{AR1}'
    put(wa, f"{c}{ST['vio']}", f'=SUM({rng})/1000', BOLD, '#,##0.0', fl=PINK)
    put(wa, f"{c}{ST['sh7']}", f'=SUMIF($A${AR0}:$A${AR1},">=7",{rng})/SUM({rng})', BOLD, PCT, fl=PINK)
    put(wa, f"{c}{ST['sh15']}", f'=SUMIF($A${AR0}:$A${AR1},">=15",{rng})/SUM({rng})', BLK, PCT)
    put(wa, f"{c}{ST['le6']}", f'=SUMIF($A${AR0}:$A${AR1},"<=6",{rng})/1000', BLK, '#,##0.0')
    put(wa, f"{c}{ST['avg']}", f'=SUMPRODUCT($B${AR0}:$B${AR1},{rng})/SUM({rng})', BLK, '0.0')
rules(wa, AR0, AR1, 3, 2 + NY, header_row_=4); rules(wa, ST['vio'], ST['avg'], 3, 2 + NY)
widths(wa, {'A': 30, 'B': 10})
for c in YC: wa.column_dimensions[c].width = 9
FA_YEARS = f"FleetByAge!$C$4:${YC[-1]}$4"; FA_MAT = f"FleetByAge!$C${AR0}:${YC[-1]}${AR1}"; FA_AGES = f"FleetByAge!$A${AR0}:$A${AR1}"
def fa_col(c): return f"FleetByAge!{c}${AR0}:{c}${AR1}"

# ================================================================ Curves
wc = wb.create_sheet('Curves')
QR0 = 6; QR1 = QR0 + NA - 1
skin(wc, 'Age curves — R(age) relative insured-claim frequency, P(age) total-loss propensity, miles by age', 'single-year resolution, ages 0–45',
     'R(a) and P(a) are NOT measured data — no public source publishes claim frequency or total-loss rate by single year of age. They are curves with 5 fitted parameters (Inputs B15, B18:B21; fitted on Calibration) chosen so the model reproduces the eight CCC claim-mix statistics on Data_CCC_Targets, given the fleet by age. Flags feed the SUMPRODUCTs on TLF_Roll. Miles (Data_EPA_Miles) are used only AFTER the fit, to split R into an exposure part and a coverage/filing part (columns I–K, docs/AGE_CURVES.md §3.4). Measured R would come from HLDI’s claim frequency by vehicle age — manual pull list, item 1.', 11, freeze='B6')
hdrs = ['Age', 'R(a)', 'P(a)', 'Miles / yr, cars', 'Miles / yr, light trucks', 'Flag age ≥ 7', 'Flag age ≤ 3', '', 'Cars on road 2024 (thousands)', 'Light trucks on road 2024 (thousands)', 'Miles / yr, fleet-weighted 2024']
header_row(wc, 4, 1, 11, hdrs, height=42)
for j, u in enumerate(['years', 'x', 'share', 'miles', 'miles', '0/1', '0/1', '', 'thousands', 'thousands', 'miles']): put(wc, f'{L(j+1)}5', u, NOTE, fl=GREY, align='center')
for i, a in enumerate(AGES):
    r = QR0 + i; put(wc, f'A{r}', a, BLUE, YR, align='center')
    put(wc, f'B{r}', f'=IF($A{r}=0,{IN["expo0"]},1)*EXP(-{IN["slope"]}*MAX($A{r}-{IN["kink"]},0))', GRN, DEC3)
    put(wc, f'C{r}', f'={IN["pmin"]}+({IN["pmax"]}-{IN["pmin"]})/(1+EXP(-($A{r}-{IN["pmid"]})/{IN["pwid"]}))', GRN, DEC3)
    put(wc, f'D{r}', f'=INDEX(Data_EPA_Miles!$B$5:$B$35,MIN($A{r},30)+1)', GRN, '#,##0')
    put(wc, f'E{r}', f'=INDEX(Data_EPA_Miles!$C$5:$C$35,MIN($A{r},30)+1)', GRN, '#,##0')
    put(wc, f'F{r}', f'=IF($A{r}>=7,1,0)', BLK, '0', align='center'); put(wc, f'G{r}', f'=IF($A{r}<=3,1,0)', BLK, '0', align='center')
    put(wc, f'I{r}', f'=IF(2024-$A{r}<{MY0},0,INDEX({SALES_C},MATCH(2024-$A{r},{SALES_Y},0))*INDEX({SV_CARS},$A{r}+1,MATCH(2024,{SV_YEARS},0)))', GRN, '#,##0')
    put(wc, f'J{r}', f'=IF(2024-$A{r}<{MY0},0,INDEX({SALES_L},MATCH(2024-$A{r},{SALES_Y},0))*INDEX({SV_LT},$A{r}+1,MATCH(2024,{SV_YEARS},0)))', GRN, '#,##0')
    put(wc, f'K{r}', f'=IF(I{r}+J{r}=0,D{r},(D{r}*I{r}+E{r}*J{r})/(I{r}+J{r}))', BLK, '#,##0')
rules(wc, QR0, QR1, 1, 7, header_row_=4); rules(wc, QR0, QR1, 9, 11, header_row_=4)
CV = dict(R=f"Curves!$B${QR0}:$B${QR1}", P=f"Curves!$C${QR0}:$C${QR1}", F7=f"Curves!$F${QR0}:$F${QR1}", LE3=f"Curves!$G${QR0}:$G${QR1}", AGE=f"Curves!$A${QR0}:$A${QR1}")
m0 = QR1 + 2
bar(wc, m0, 1, 4, 'R decomposition, vehicles 7+ vs 0–6 (2024 fleet weights)', BLUEF, WHITE_B, 16)
N24 = fa_col(YC[YEARS.index(2024)])
for k, (lab_, f_, note) in enumerate([
    ('Fleet-weighted miles ratio 7+ / 0–6', f'=(SUMPRODUCT($K${QR0}:$K${QR1},$F${QR0}:$F${QR1},{N24})/SUMPRODUCT($F${QR0}:$F${QR1},{N24}))/(SUMPRODUCT($K${QR0}:$K${QR1},1-$F${QR0}:$F${QR1},{N24})/SUMPRODUCT(1-$F${QR0}:$F${QR1},{N24}))', 'EPA T3.14 miles × 2024 fleet (≈0.65)'),
    ('Fitted R ratio 7+ / 0–6 (claims per vehicle)', f'=(SUMPRODUCT($B${QR0}:$B${QR1},$F${QR0}:$F${QR1},{N24})/SUMPRODUCT($F${QR0}:$F${QR1},{N24}))/(SUMPRODUCT($B${QR0}:$B${QR1},1-$F${QR0}:$F${QR1},{N24})/SUMPRODUCT(1-$F${QR0}:$F${QR1},{N24}))', 'claims per vehicle, 7+ vs 0–6 (≈0.56)'),
    ('Residual = coverage / filing', f'=B{m0+2}/B{m0+1}', 'liability-only, higher deductibles, unfiled small claims on older cars (≈0.87)')]):
    r = m0 + 1 + k; put(wc, f'A{r}', lab_, BOLD if k == 2 else BLK, fl=(PINK if k == 2 else None)); put(wc, f'B{r}', f_, BOLD if k == 2 else BLK, DEC3, fl=(PINK if k == 2 else None)); put(wc, f'C{r}', note, NOTE)
rules(wc, m0 + 1, m0 + 3, 1, 2)
widths(wc, {'A': 34, 'B': 10, 'C': 10, 'D': 12, 'E': 12, 'F': 9, 'G': 9, 'H': 3, 'I': 14, 'J': 14, 'K': 14})

# ================================================================ Calibration (Solver)
wk = wb.create_sheet('Calibration')
KR0 = 13; KR1 = KR0 + NA - 1
skin(wk, 'Calibration — fit R(age) and P(age) to CCC 2024 with Solver', 'five knobs, eight gauges, one score',
     'Working parameters B4:B8 (pink) are what Solver changes; they start at the fitted values. After a run, paste B4 → Inputs B15 and B5:B8 → Inputs B18:B21 so the model does not depend on Solver being re-run. The 2024 fleet column links from FleetByAge.', 15, freeze='A13')
for r, lab_, v, fmt, b in [(4, 'R decline per year after the kink', FIT['r_slope'], DEC4, '0 to 0.5'), (5, 'P floor (pmin)', FIT['pmin'], DEC3, '0.01 to 0.3'), (6, 'P ceiling (pmax)', FIT['pmax'], DEC3, '0.3 to 0.95, and ≥ B5'),
                            (7, 'P midpoint age', FIT['pmid'], '0.00', '0 to 30'), (8, 'P width', FIT['pwid'], '0.00', '0.3 to 12')]:
    put(wk, f'A{r}', lab_, BLK); put(wk, f'B{r}', v, BLUE, fmt, fl=PINK, align='right'); put(wk, f'C{r}', b, NOTE)
put(wk, 'A9', 'Exposure age 0 (Inputs)', BLK); put(wk, 'B9', f'={IN["expo0"]}', GRN, '0.00', align='right'); put(wk, 'A10', 'Kink age (Inputs)', BLK); put(wk, 'B10', f'={IN["kink"]}', GRN, YR, align='right')
rules(wk, 4, 10, 1, 3)
header_row(wk, 12, 1, 7, ['Age', 'Fleet 2024 (thousands)', 'R(a)', 'P(a)', 'Claims = fleet × R', 'Total losses = claims × P', 'Repairables = claims − TL'], height=42)
for i, a in enumerate(AGES):
    r = KR0 + i; put(wk, f'A{r}', a, BLUE, YR, align='center')
    put(wk, f'B{r}', f'=INDEX({FA_MAT},MATCH($A{r},{FA_AGES},0),MATCH(2024,{FA_YEARS},0))', GRN, '#,##0')
    put(wk, f'C{r}', f'=IF($A{r}=0,$B$9,1)*EXP(-$B$4*MAX($A{r}-$B$10,0))', BLK, DEC3)
    put(wk, f'D{r}', f'=$B$5+($B$6-$B$5)/(1+EXP(-($A{r}-$B$7)/$B$8))', BLK, DEC3)
    put(wk, f'E{r}', f'=B{r}*C{r}', BLK, '#,##0'); put(wk, f'F{r}', f'=E{r}*D{r}', BLK, '#,##0'); put(wk, f'G{r}', f'=E{r}-F{r}', BLK, '#,##0')
rules(wk, KR0, KR1, 1, 7, header_row_=12)
A_, E_, F_, G_ = f'$A${KR0}:$A${KR1}', f'$E${KR0}:$E${KR1}', f'$F${KR0}:$F${KR1}', f'$G${KR0}:$G${KR1}'
stats = [
 ('tlf',    'Total-loss frequency',                 f'=SUM({F_})/SUM({E_})', PCT),
 ('tl7',    'Total losses from vehicles 7+',        f'=SUMIF({A_},">=7",{F_})/SUM({F_})', PCT),
 ('rp7',    'Repairables from vehicles 7+',         f'=SUMIF({A_},">=7",{G_})/SUM({G_})', PCT),
 ('rp3',    'Repairables from vehicles ≤3',         f'=SUMIF({A_},"<=3",{G_})/SUM({G_})', PCT),
 ('age_cl', 'Average age, claim vehicles',          f'=SUMPRODUCT({A_},{E_})/SUM({E_})', '0.00'),
 ('age_rp', 'Average age, repairables',             f'=SUMPRODUCT({A_},{G_})/SUM({G_})', '0.00'),
 ('age_tl', 'Average age, total losses',            f'=SUMPRODUCT({A_},{F_})/SUM({F_})', '0.00'),
 ('p03',    'Total-loss rate, vehicles ≤3',         f'=SUMIF({A_},"<=3",{F_})/SUMIF({A_},"<=3",{E_})', PCT),
]
header_row(wk, 12, 9, 15, ['Key', 'Statistic (2024)', 'Model', 'CCC target', 'Tolerance', 'Scaled sq. error', 'Source quote'], height=42, center=False)
for i, (k, lab_, f_, fmt) in enumerate(stats):
    r = 13 + i; trow = TR0 + [t[2] for t in TGT if t[0] == 'FIT'].index(k)
    put(wk, f'I{r}', k, BLK); put(wk, f'J{r}', lab_, BLK); put(wk, f'K{r}', f_, BLK, fmt)
    put(wk, f'L{r}', f'=Data_CCC_Targets!$E${trow}', GRN, fmt); put(wk, f'M{r}', f'=Data_CCC_Targets!$F${trow}', GRN, ('0.00' if k.startswith('age') else PCT2))
    put(wk, f'N{r}', f'=((K{r}-L{r})/M{r})^2', BLK, '0.000'); put(wk, f'O{r}', f'=Data_CCC_Targets!$G${trow}', GRN)
rules(wk, 13, 20, 9, 15, header_row_=12)
for c_ in 'JKLM': wk[f'{c_}22'].fill = PINK
put(wk, 'J22', 'LOSS — Solver objective, minimise', BOLD, fl=PINK); put(wk, 'N22', '=SUM(N13:N20)', BOLD, '0.000', fl=PINK)
put(wk, 'J23', 'Reference: the Python fit (scripts/age_curves.py) lands at ~1.0. One unit = one statistic off by exactly its tolerance.', NOTE)
bar(wk, 25, 9, 15, 'SOLVER SET-UP (Excel → Data → Solver; on Mac enable once via Tools → Excel Add-ins → Solver)', BLUEF, WHITE_B, 16)
solver = ["1. Set Objective: $N$22, To: Min.",
          "2. By Changing Variable Cells: $B$4:$B$8.",
          "3. Constraints: B4 ≥ 0; B4 ≤ 0.5; B5 ≥ 0.01; B5 ≤ 0.3; B6 ≥ 0.3; B6 ≤ 0.95; B6 ≥ B5; B7 ≥ 0; B7 ≤ 30; B8 ≥ 0.3; B8 ≤ 12.",
          "4. Method: GRG Nonlinear. Options → GRG Nonlinear tab: tick Use Multistart, Population Size 100, Random Seed 7, tick Require Bounds on Variables. Untick 'Make Unconstrained Variables Non-Negative'.",
          "5. Solve. Expect LOSS ≈ 1.0 and B4:B8 within a few thousandths of the starting values. Paste values: B4 → Inputs!B15; B5:B8 → Inputs!B18:B21.",
          "In words: turn the five knobs until the model’s eight 2024 statistics match CCC’s published ones. Multistart = start from many random knob positions so Solver does not settle in a shallow dip."]
for i, t in enumerate(solver): put(wk, f'I{26+i}', t, BLK)
widths(wk, {'A': 34, 'B': 12, 'C': 22, 'D': 9, 'E': 15, 'F': 18, 'G': 18, 'H': 3, 'I': 8, 'J': 36, 'K': 10, 'L': 11, 'M': 10, 'N': 13, 'O': 90})

# ================================================================ TLF_Roll
wr = wb.create_sheet('TLF_Roll')
skin(wr, 'Baseline total-loss frequency from demographics alone, by year — and the out-of-sample tests', 'R and P frozen; only the fleet changes',
     'claims index = Σ_a fleet(a,t)·R(a); TL index = Σ_a fleet(a,t)·R(a)·P(a); baseline TLF = TL ÷ claims. The YoY row is the demographic contribution to TLF; everything else in actual TLF is level (spread cycle, technology, filing behaviour). Checks below compare with CCC for years never used in the fit.', 2 + NY, freeze='C6')
put(wr, 'A4', 'Line item  /  calendar year →', WHITE_B, fl=NAVY); put(wr, 'B4', 'units', WHITE_B, fl=NAVY, align='center')
year_header(wr, 4, YC, YEARS, link_from='inputs')
RW = {'cl': 6, 'tl': 7, 'tlf': 8, 'drift': 9, 'cl7': 10, 'tl7': 11, 'rp7': 12, 'rp3': 13, 'age_cl': 14, 'age_tl': 15, 'age_rp': 16, 'p03': 17, 'vio': 18, 'sh7': 19, 'avg': 20, 'tlgr': 21}
lab = {'cl': ('Claims index  Σ fleet · R', 'thousands'), 'tl': ('Total-loss index  Σ fleet · R · P', 'thousands'), 'tlf': ('Baseline total-loss frequency (demographics only)', '%'), 'drift': ('   YoY change in baseline TLF', 'pp'),
       'cl7': ('Claims share, vehicles 7+', '%'), 'tl7': ('Total-loss share, vehicles 7+', '%'), 'rp7': ('Repairable share, vehicles 7+', '%'), 'rp3': ('Repairable share, vehicles ≤3', '%'),
       'age_cl': ('Average age, claim vehicles', 'years'), 'age_tl': ('Average age, total losses', 'years'), 'age_rp': ('Average age, repairables', 'years'), 'p03': ('Total-loss rate, vehicles ≤3', '%'),
       'vio': ('Vehicles in operation (FleetByAge)', 'millions'), 'sh7': ('Fleet share aged 7+ (FleetByAge)', '%'), 'avg': ('Average fleet age (t − MY + 0.5)', 'years'), 'tlgr': ('   Total-loss index YoY (demographic TL pool growth)', '%')}
put(wr, 'A5', 'Fleet roll → claims → total losses', BOLD, fl=GREY); put(wr, 'B5', '', fl=GREY)
for c in YC: wr[f'{c}5'].fill = GREY
for k, r in RW.items():
    key = k in ('tlf', 'drift'); put(wr, f'A{r}', lab[k][0], BOLD if key else BLK, fl=(PINK if key else None)); put(wr, f'B{r}', lab[k][1], NOTE, align='center', fl=(PINK if key else None))
for i, c in enumerate(YC):
    F = fa_col(c); prev = YC[i - 1] if i > 0 else None
    put(wr, f'{c}{RW["cl"]}', f'=SUMPRODUCT({F},{CV["R"]})', BLK, '#,##0')
    put(wr, f'{c}{RW["tl"]}', f'=SUMPRODUCT({F},{CV["R"]},{CV["P"]})', BLK, '#,##0')
    put(wr, f'{c}{RW["tlf"]}', f'={c}{RW["tl"]}/{c}{RW["cl"]}', BOLD, PCT2, fl=PINK)
    put(wr, f'{c}{RW["drift"]}', (f'=({c}{RW["tlf"]}-{prev}{RW["tlf"]})*100' if prev else ''), BOLD, '+0.00;-0.00', fl=PINK)
    put(wr, f'{c}{RW["cl7"]}', f'=SUMPRODUCT({F},{CV["R"]},{CV["F7"]})/{c}{RW["cl"]}', BLK, PCT)
    put(wr, f'{c}{RW["tl7"]}', f'=SUMPRODUCT({F},{CV["R"]},{CV["P"]},{CV["F7"]})/{c}{RW["tl"]}', BLK, PCT)
    put(wr, f'{c}{RW["rp7"]}', f'=(SUMPRODUCT({F},{CV["R"]},{CV["F7"]})-SUMPRODUCT({F},{CV["R"]},{CV["P"]},{CV["F7"]}))/({c}{RW["cl"]}-{c}{RW["tl"]})', BLK, PCT)
    put(wr, f'{c}{RW["rp3"]}', f'=(SUMPRODUCT({F},{CV["R"]},{CV["LE3"]})-SUMPRODUCT({F},{CV["R"]},{CV["P"]},{CV["LE3"]}))/({c}{RW["cl"]}-{c}{RW["tl"]})', BLK, PCT)
    put(wr, f'{c}{RW["age_cl"]}', f'=SUMPRODUCT({F},{CV["R"]},{CV["AGE"]})/{c}{RW["cl"]}', BLK, '0.0')
    put(wr, f'{c}{RW["age_tl"]}', f'=SUMPRODUCT({F},{CV["R"]},{CV["P"]},{CV["AGE"]})/{c}{RW["tl"]}', BLK, '0.0')
    put(wr, f'{c}{RW["age_rp"]}', f'=(SUMPRODUCT({F},{CV["R"]},{CV["AGE"]})-SUMPRODUCT({F},{CV["R"]},{CV["P"]},{CV["AGE"]}))/({c}{RW["cl"]}-{c}{RW["tl"]})', BLK, '0.0')
    put(wr, f'{c}{RW["p03"]}', f'=SUMPRODUCT({F},{CV["R"]},{CV["P"]},{CV["LE3"]})/SUMPRODUCT({F},{CV["R"]},{CV["LE3"]})', BLK, PCT)
    put(wr, f'{c}{RW["vio"]}', f"=FleetByAge!{c}{ST['vio']}", GRN, '#,##0.0'); put(wr, f'{c}{RW["sh7"]}', f"=FleetByAge!{c}{ST['sh7']}", GRN, PCT); put(wr, f'{c}{RW["avg"]}', f"=FleetByAge!{c}{ST['avg']}", GRN, '0.0')
    put(wr, f'{c}{RW["tlgr"]}', (f'=({c}{RW["tl"]}/{prev}{RW["tl"]}-1)*100' if prev else ''), BLK, '+0.0;-0.0')
rules(wr, 6, 21, 3, 2 + NY, header_row_=4)
CK0 = 24
bar(wr, CK0, 1, 9, 'Checks against CCC — R and P frozen at the 2024 fit; rows marked OOS were never used in the fit', BLUEF, WHITE_B, 16)
header_row(wr, CK0 + 1, 1, 9, ['Use', 'Year', 'Statistic', 'CCC actual', 'Model', 'Error', 'Quote'], wrap=False, center=False)
wr.merge_cells(start_row=CK0 + 1, start_column=7, end_row=CK0 + 1, end_column=9)
keymap = {'tlf': 'tlf', 'tl7': 'tl7', 'rp7': 'rp7', 'rp3': 'rp3', 'age_cl': 'age_cl', 'age_rp': 'age_rp', 'age_tl': 'age_tl', 'p03': 'p03', 'vio': 'vio', 'sh7': 'sh7'}
r = CK0 + 2
for use, year, key, statname, val, tol, quote, page in TGT:
    if key not in keymap: continue
    isage = key.startswith('age')
    put(wr, f'A{r}', use, BLK, align='center', fl=(PINK if use == 'OOS' else None)); put(wr, f'B{r}', year, BLUE, YR, align='center'); put(wr, f'C{r}', statname, BLK)
    put(wr, f'D{r}', val, BLUE, ('0.0' if isage else ('#,##0' if key == 'vio' else PCT)))
    put(wr, f'E{r}', f'=INDEX($C${RW[keymap[key]]}:${YC[-1]}${RW[keymap[key]]},1,MATCH(B{r},$C$4:${YC[-1]}$4,0))', BLK, ('0.0' if isage else ('#,##0.0' if key == 'vio' else PCT)))
    put(wr, f'F{r}', (f'=E{r}-D{r}' if (isage or key == 'vio') else f'=(E{r}-D{r})*100'), BLK, ('+0.0;-0.0' if (isage or key == 'vio') else '+0.0"pp";-0.0"pp"'))
    put(wr, f'G{r}', quote, NOTE); wr.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9); r += 1
rules(wr, CK0 + 2, r - 1, 1, 6, header_row_=CK0 + 1)
c19, c25 = YC[YEARS.index(2019)], YC[YEARS.index(2025)]
for c_ in 'ABCDEF': wr[f'{c_}{r+1}'].fill = PINK
put(wr, f'A{r+1}', 'Demographic drift 2019→2025, pp per year', BOLD, fl=PINK); put(wr, f'C{r+1}', 'model vs actual CCC rise (+0.65pp/yr, 19.2% → 23.1%)', BLK, fl=PINK); put(wr, f'D{r+1}', '=(0.231-0.192)*100/6', BLK, '+0.00', fl=PINK)
put(wr, f'E{r+1}', f'=({c25}{RW["tlf"]}-{c19}{RW["tlf"]})*100/6', BOLD, '+0.00', fl=PINK); put(wr, f'F{r+1}', f'=E{r+1}/D{r+1}', BOLD, PCT, fl=PINK); put(wr, f'G{r+1}', 'share of the 2019→2025 rise that is demographics (~24%); the rest is level', NOTE)
widths(wr, {'A': 50, 'B': 10})
for c in YC: wr.column_dimensions[c].width = 9

# ================================================================ Spread_Reg
wsp = wb.create_sheet('Spread_Reg')
skin(wsp, 'Totaling-spread regression — ΔTLF (pp, YoY) = a + b × spread(t−1), CCC quarterly', 'the cycle term for total-loss frequency',
     'spread = repair CPI YoY − used-car CPI YoY (BLS NSA, quarterly mean of monthly values; Data_Spread). ΔTLF = CCC quarterly TLF minus the same quarter a year earlier (Data_CCC_TLF_Quarterly). One-quarter lag. Reproduces scripts/analysis_20260911.py: all-loss a = 0.599, b = 0.0815, R² 0.81, n = 27; non-comp a = 0.609, b = 0.0834, R² 0.79.', 22, freeze='A5')
_, _, sp_rows = read_csv('totaling_spread_quarterly.csv'); sp_rows = [r for r in sp_rows if r[0] >= '2016Q1']
_, _, tq_rows = read_csv('ccc_tlf_quarterly.csv')
header_row(wsp, 4, 1, 4, ['Quarter', 'Spread (pp)', 'Months', 'Note'], wrap=False)
for i, r in enumerate(sp_rows):
    put(wsp, f'A{5+i}', r[0], BLUE, align='center'); put(wsp, f'B{5+i}', float(r[1]), BLUE, '0.00'); put(wsp, f'C{5+i}', int(r[2]), BLUE, '0', align='center'); put(wsp, f'D{5+i}', r[3], NOTE)
SP0, SP1 = 5, 4 + len(sp_rows); rules(wsp, SP0, SP1, 1, 4, header_row_=4)
header_row(wsp, 4, 6, 8, ['Quarter', 'TLF non-comp %', 'TLF all-loss %'])
for i, r in enumerate(tq_rows):
    put(wsp, f'F{5+i}', r[0], BLUE, align='center'); put(wsp, f'G{5+i}', float(r[1]), BLUE, '0.0'); put(wsp, f'H{5+i}', float(r[2]), BLUE, '0.0')
TQ0, TQ1 = 5, 4 + len(tq_rows); rules(wsp, TQ0, TQ1, 6, 8, header_row_=4)
pairs_q = [r[0] for r in tq_rows if any(t[0] == f"{int(r[0][:4])-1}{r[0][4:]}" for t in tq_rows)]
header_row(wsp, 4, 10, 17, ['Quarter t', 'Year', 'Q', 'Prior-year quarter', 'Lag quarter t−1', 'ΔTLF all-loss (pp)', 'ΔTLF non-comp (pp)', 'Spread (t−1)'], height=30)
for i, q in enumerate(pairs_q):
    r = 5 + i
    put(wsp, f'J{r}', q, BLUE, align='center'); put(wsp, f'K{r}', f'=VALUE(LEFT(J{r},4))', BLK, '0'); put(wsp, f'L{r}', f'=VALUE(RIGHT(J{r},1))', BLK, '0', align='center')
    put(wsp, f'M{r}', f'=(K{r}-1)&"Q"&L{r}', BLK, align='center'); put(wsp, f'N{r}', f'=IF(L{r}=1,(K{r}-1)&"Q4",K{r}&"Q"&(L{r}-1))', BLK, align='center')
    put(wsp, f'O{r}', f'=INDEX($H${TQ0}:$H${TQ1},MATCH(J{r},$F${TQ0}:$F${TQ1},0))-INDEX($H${TQ0}:$H${TQ1},MATCH(M{r},$F${TQ0}:$F${TQ1},0))', BLK, '+0.0;-0.0')
    put(wsp, f'P{r}', f'=INDEX($G${TQ0}:$G${TQ1},MATCH(J{r},$F${TQ0}:$F${TQ1},0))-INDEX($G${TQ0}:$G${TQ1},MATCH(M{r},$F${TQ0}:$F${TQ1},0))', BLK, '+0.0;-0.0')
    put(wsp, f'Q{r}', f'=INDEX($B${SP0}:$B${SP1},MATCH(N{r},$A${SP0}:$A${SP1},0))', BLK, '0.00')
PR0, PR1 = 5, 4 + len(pairs_q); rules(wsp, PR0, PR1, 10, 17, header_row_=4)
header_row(wsp, 4, 19, 21, ['Regression', 'All-loss', 'Non-comp'], wrap=False, center=False)
regrows = [('Slope b — pp of TLF per pp of spread', 'SLOPE', DEC4), ('Intercept a — pp per year, the unexplained drift', 'INTERCEPT', DEC3), ('R²', 'RSQ', DEC3), ('n', 'COUNT', '0')]
for i, (lab_, fn, fmt) in enumerate(regrows):
    r = 5 + i; hl = PINK if fn in ('SLOPE', 'INTERCEPT') else None; put(wsp, f'S{r}', lab_, BOLD if hl else BLK, fl=hl)
    if fn == 'COUNT': put(wsp, f'T{r}', f'=COUNT($O${PR0}:$O${PR1})', BLK, fmt); put(wsp, f'U{r}', f'=COUNT($P${PR0}:$P${PR1})', BLK, fmt)
    else:
        put(wsp, f'T{r}', f'={fn}($O${PR0}:$O${PR1},$Q${PR0}:$Q${PR1})', BOLD if hl else BLK, fmt, fl=hl); put(wsp, f'U{r}', f'={fn}($P${PR0}:$P${PR1},$Q${PR0}:$Q${PR1})', BOLD if hl else BLK, fmt, fl=hl)
rules(wsp, 5, 8, 19, 21, header_row_=4)
bar(wsp, 10, 19, 21, 'Prediction', BLUEF, WHITE_B, 16)
put(wsp, 'S11', 'Spread quarter to use (t−1)', BLK); put(wsp, 'T11', '2026Q1', BLUE, align='center'); put(wsp, 'S12', 'spread(t−1)', BLK); put(wsp, 'T12', f'=INDEX($B${SP0}:$B${SP1},MATCH(T11,$A${SP0}:$A${SP1},0))', BLK, '0.00')
put(wsp, 'S13', 'Predicted ΔTLF all-loss, quarter t (pp YoY)', BOLD, fl=PINK); put(wsp, 'T13', '=T6+T5*T12', BOLD, '+0.00', fl=PINK)
put(wsp, 'S14', 'Live test: mgmt-cited CCC 2Q26 23.3 vs 22.4 (2026-09-10 call)', NOTE); put(wsp, 'T14', 0.9, BLUE, '+0.00'); put(wsp, 'S15', 'Error (pp)', BLK); put(wsp, 'T15', '=T13-T14', BLK, '+0.00')
rules(wsp, 11, 15, 19, 20)
put(wsp, 'S17', 'Why the intercept matters: with spread = 0 the model still adds ~0.6pp/yr of TLF; the cohort roll (TLF_Roll) explains ~0.16pp of that. Use the spread term for 12 months, not for the decade. MODEL_BLUEPRINT §3.', NOTE, wrap=True); wsp.merge_cells('S17:V19')
widths(wsp, {'A': 9, 'B': 11, 'C': 8, 'D': 36, 'E': 2, 'F': 9, 'G': 13, 'H': 13, 'I': 2, 'J': 10, 'K': 6, 'L': 5, 'M': 16, 'N': 14, 'O': 16, 'P': 16, 'Q': 11, 'R': 2, 'S': 46, 'T': 11, 'U': 11})

# ================================================================ RPU_Reg
wq = wb.create_sheet('RPU_Reg')
skin(wq, 'RPU chain — ASP ~ used-car CPI (nowcast) and service RPU ~ ASP (elasticity)', 'the price side of revenue per unit',
     'Step 1: Copart US insurance ASP YoY (transcripts, Data_ReportedUnits) on used-car CPI YoY (BLS CUSR0000SETA02, SA, fiscal-quarter average of the three months vs the same months a year earlier; Data_CPI col B). Step 2: implied service RPU YoY on global ASP YoY, n = 17 (Data_Elasticity). Reference values: step 1 a = 3.23, b = 0.599, r = 0.77 (n = 15); step 2 b = 0.514, a = 4.13, R² 0.61; total-RPU b = 0.752, a = 2.84.', 16, freeze='A5')
FQM = [("FY2023 Q1", "2022-08", "2022-09", "2022-10"), ("FY2023 Q2", "2022-11", "2022-12", "2023-01"), ("FY2023 Q3", "2023-02", "2023-03", "2023-04"), ("FY2023 Q4", "2023-05", "2023-06", "2023-07"),
       ("FY2024 Q1", "2023-08", "2023-09", "2023-10"), ("FY2024 Q2", "2023-11", "2023-12", "2024-01"), ("FY2024 Q3", "2024-02", "2024-03", "2024-04"), ("FY2024 Q4", "2024-05", "2024-06", "2024-07"),
       ("FY2025 Q1", "2024-08", "2024-09", "2024-10"), ("FY2025 Q2", "2024-11", "2024-12", "2025-01"), ("FY2025 Q3", "2025-02", "2025-03", "2025-04"), ("FY2025 Q4", "2025-05", "2025-06", "2025-07"),
       ("FY2026 Q1", "2025-08", "2025-09", "2025-10"), ("FY2026 Q2", "2025-11", "2025-12", "2026-01"), ("FY2026 Q3", "2026-02", "2026-03", "2026-04"), ("FY2026 Q4", "2026-05", "2026-06", "2026-07")]
_, cpi_hdr, cpi_rows = read_csv('cprt_cpi_three_series.csv'); CPI0, CPI1 = 5, 4 + len(cpi_rows)
_, ru_hdr, ru_rows = read_csv('reported_units.csv'); RU0, RU1 = 5, 4 + len(ru_rows)
header_row(wq, 4, 1, 8, ['Fiscal quarter', 'Month 1', 'Month 2', 'Month 3', 'CPI used cars, quarter avg', 'Same quarter, prior year', 'CPI YoY %', 'US insurance ASP YoY % (reported)'], height=42)
cpiA = f'Data_CPI!$A${CPI0}:$A${CPI1}'; cpiB = f'Data_CPI!$B${CPI0}:$B${CPI1}'
for i, (fq, m1, m2, m3) in enumerate(FQM):
    r = 5 + i; put(wq, f'A{r}', fq, BLUE); put(wq, f'B{r}', m1, BLUE, align='center'); put(wq, f'C{r}', m2, BLUE, align='center'); put(wq, f'D{r}', m3, BLUE, align='center')
    cur = ",".join(f'INDEX({cpiB},MATCH({c}{r},{cpiA},0))' for c in 'BCD'); pri = ",".join(f'INDEX({cpiB},MATCH((VALUE(LEFT({c}{r},4))-1)&RIGHT({c}{r},3),{cpiA},0))' for c in 'BCD')
    put(wq, f'E{r}', f'=IFERROR(AVERAGE({cur}),"")', GRN, '0.0'); put(wq, f'F{r}', f'=IFERROR(AVERAGE({pri}),"")', GRN, '0.0')
    put(wq, f'G{r}', f'=IF(OR(E{r}="",F{r}=""),"",(E{r}/F{r}-1)*100)', BLK, '+0.0;-0.0')
    put(wq, f'H{r}', f'=IFERROR(INDEX(Data_ReportedUnits!$F${RU0}:$F${RU1},MATCH(A{r},Data_ReportedUnits!$A${RU0}:$A${RU1},0)),"")', GRN, '+0.0;-0.0')
Q0, Q1 = 5, 4 + 15; rules(wq, 5, 4 + len(FQM), 1, 8, header_row_=4)
for c_ in 'ABCDEFGH': wq[f'{c_}{Q1+1}'].fill = GREY
header_row(wq, 4, 10, 12, ['Step 1: ASP = a + b × used-car CPI', 'Value', ''], wrap=False, center=False)
put(wq, 'J5', 'Slope b', BOLD, fl=PINK); put(wq, 'K5', f'=SLOPE(H{Q0}:H{Q1},G{Q0}:G{Q1})', BOLD, DEC3, fl=PINK)
put(wq, 'J6', 'Intercept a (pp)', BOLD, fl=PINK); put(wq, 'K6', f'=INTERCEPT(H{Q0}:H{Q1},G{Q0}:G{Q1})', BOLD, '0.00', fl=PINK)
put(wq, 'J7', 'Correlation r', BLK); put(wq, 'K7', f'=CORREL(H{Q0}:H{Q1},G{Q0}:G{Q1})', BLK, DEC3); put(wq, 'J8', 'n', BLK); put(wq, 'K8', f'=COUNT(H{Q0}:H{Q1})', BLK, '0')
rules(wq, 5, 8, 10, 11, header_row_=4)
put(wq, 'J9', 'Grey row: FY2026 Q4 has CPI but the reported ASP is not in the CSV yet (call: US ins ASP +3.7%). Extend the regression rows when the CSV is updated.', NOTE)
_, el_hdr, el_rows = read_csv('elasticity_rebuild.csv'); EL0, EL1 = 5, 4 + len(el_rows)
elE = f'Data_Elasticity!$E${EL0}:$E${EL1}'; elF = f'Data_Elasticity!$F${EL0}:$F${EL1}'; elG = f'Data_Elasticity!$G${EL0}:$G${EL1}'
header_row(wq, 11, 10, 12, ['Step 2: RPU = a + b × ASP  (global, n = 17)', 'Service RPU', 'Total RPU (mgmt def.)'], wrap=False, center=False)
put(wq, 'J12', 'Slope b — elasticity', BOLD, fl=PINK); put(wq, 'K12', f'=SLOPE({elF},{elE})', BOLD, DEC3, fl=PINK); put(wq, 'L12', f'=SLOPE({elG},{elE})', BOLD, DEC3, fl=PINK)
put(wq, 'J13', 'Intercept a — fee/mix growth at flat ASP (pp)', BOLD, fl=PINK); put(wq, 'K13', f'=INTERCEPT({elF},{elE})', BOLD, '0.00', fl=PINK); put(wq, 'L13', f'=INTERCEPT({elG},{elE})', BOLD, '0.00', fl=PINK)
put(wq, 'J14', 'R²', BLK); put(wq, 'K14', f'=RSQ({elF},{elE})', GRN, DEC3); put(wq, 'L14', f'=RSQ({elG},{elE})', GRN, DEC3)
put(wq, 'J15', 'n', BLK); put(wq, 'K15', f'=COUNT({elF})', GRN, '0'); put(wq, 'L15', f'=COUNT({elG})', GRN, '0')
rules(wq, 12, 15, 10, 12, header_row_=11)
put(wq, 'J16', 'Caveat: residual autocorrelation → effective n ≈ 6; 95% CI on the service-RPU slope [0.29, 0.73]. The INTERCEPT is the robust finding (fee/mix adds ~4pp/yr at flat ASP), and it decelerated in FY26 (+1.3 to +3.5pp).', NOTE, wrap=True); wq.merge_cells('J16:P17'); wq.row_dimensions[16].height = 30
bar(wq, 19, 10, 12, 'Chain — type a used-car CPI YoY assumption', BLUEF, WHITE_B, 16)
put(wq, 'J20', 'Used-car CPI YoY % (assumption)', BLK); put(wq, 'K20', -2.0, BLUE, '+0.0;-0.0')
put(wq, 'J21', '→ US insurance ASP YoY %', BLK); put(wq, 'K21', '=K6+K5*K20', BLK, '+0.0;-0.0')
put(wq, 'J22', '→ Service RPU YoY %', BOLD, fl=PINK); put(wq, 'K22', '=K13+K12*K21', BOLD, '+0.0;-0.0', fl=PINK)
put(wq, 'J23', '→ Total RPU YoY % (mgmt definition)', BLK); put(wq, 'K23', '=L13+L12*K21', BLK, '+0.0;-0.0')
rules(wq, 20, 23, 10, 11)
put(wq, 'J24', 'FY26Q4 check (IN-SAMPLE, not a back-test): used-car CPI −1.9% → ASP +2.1% (actual +3.5%) → service RPU +5.2% (actual implied +4.4%). The miss is mostly the ASP step. RPU is low-variance: +3.3% to +8.9% across used-car CPI −8% to +10%.', NOTE, wrap=True); wq.merge_cells('J24:P25')
widths(wq, {'A': 13, 'B': 9, 'C': 9, 'D': 9, 'E': 13, 'F': 13, 'G': 10, 'H': 16, 'I': 2, 'J': 46, 'K': 12, 'L': 18})

# ================================================================ Checks
wch = wb.create_sheet('Checks')
skin(wch, 'Checks — every row should read OK after Excel recalculates', 'reference values from the Python scripts', 'A FAIL after copying tabs into another workbook almost always means a sheet was renamed or a dependency tab was not copied. Tolerances are deliberately loose; they catch broken references, not rounding.', 5, freeze='A5')
header_row(wch, 4, 1, 5, ['Check', 'Value', 'Target', 'Tolerance', 'Status'], wrap=False, center=False)
c24 = YC[YEARS.index(2024)]; c13 = YC[0]
checks = [
 ('Calibration loss (Solver objective) small', '=Calibration!N22', 0, 3),
 ('Baseline TLF 2024 ≈ CCC 22.3%', f'=TLF_Roll!{c24}{RW["tlf"]}', 0.223, 0.005),
 ('Vehicles in operation 2024 ≈ 292M (Experian) / 289M (S&P)', f'=TLF_Roll!{c24}{RW["vio"]}', 292, 9),
 ('Fleet share aged 7+ 2024 ≈ 66%', f'=TLF_Roll!{c24}{RW["sh7"]}', 0.66, 0.025),
 ('Survival at age 0 = 1 in every year', f'=MIN(Survival!C{CR0}:{YC[-1]}{CR0},Survival!C{LR0}:{YC[-1]}{LR0})', 1, 0.0001),
 ('Fleet 2013 ≈ 2013 census light vehicles (~243M)', f'=Fleet!{c13}{TOTR}/1000', 243, 12),
 ('Spread regression slope ≈ 0.0815 (all-loss)', '=Spread_Reg!T5', 0.0815, 0.003),
 ('Spread regression intercept ≈ 0.599 (all-loss)', '=Spread_Reg!T6', 0.599, 0.03),
 ('Service-RPU elasticity ≈ 0.514', '=RPU_Reg!K12', 0.514, 0.01),
 ('Service-RPU intercept ≈ 4.13', '=RPU_Reg!K13', 4.13, 0.1),
 ('ASP ~ CPI slope ≈ 0.60 (n = 15)', '=RPU_Reg!K5', 0.599, 0.05),
 ('Total-loss rate for vehicles ≤3, 2024 ≈ 10%', f'=TLF_Roll!{c24}{RW["p03"]}', 0.10, 0.01),
]
for i, (lab_, f_, tgt, tol) in enumerate(checks):
    r = 5 + i; put(wch, f'A{r}', lab_, BLK); put(wch, f'B{r}', f_, GRN, '0.0000'); put(wch, f'C{r}', tgt, BLUE, '0.0000'); put(wch, f'D{r}', tol, BLUE, '0.0000')
    put(wch, f'E{r}', f'=IF(ABS(B{r}-C{r})<=D{r},"OK","FAIL")', BOLD, align='center')
NCK = len(checks); rules(wch, 5, 4 + NCK, 1, 5, header_row_=4)
for c_ in 'ABCDE': wch[f'{c_}{6+NCK}'].fill = PINK
put(wch, f'A{6+NCK}', 'ALL CHECKS', BOLD, fl=PINK); put(wch, f'E{6+NCK}', f'=IF(COUNTIF(E5:E{4+NCK},"OK")={NCK},"OK","FAIL")', BOLD, fl=PINK, align='center')
widths(wch, {'A': 60, 'B': 12, 'C': 12, 'D': 12, 'E': 9})

# ================================================================ README
wrm = wb.create_sheet('README')
skin(wrm, 'Read me — how this workbook is built and how to copy from it', f'built {TODAY}', 'Built by scripts/build_intermediate_xlsx.py from the committed CSVs in data/csv/. Rebuild after any data update; do not hand-edit Data_* tabs. Verified by scripts/verify_intermediate_xlsx.py (pycel) — all checks OK at build time.', 2, tab=C_YEL)
readme = [
 ('HOW TO USE', None),
 ('Copy whole sheets, not cells', 'Right-click a tab → Move or Copy → To book: CPRT Model → tick Create a copy. Copy the tabs a formula sheet depends on FIRST (Inputs, Data_Sales, Data_EPA_Survival, Data_EPA_Miles, Data_CCC_Targets, then Survival, Fleet, FleetByAge, Curves, then TLF_Roll / Calibration). References are explicit sheet names, no named ranges, so they survive the copy as long as tab names are unchanged.'),
 ('Colour legend', 'BLUE font = hard-coded input · BLACK = formula · GREEN = link to another sheet · PINK fill = calibrated parameter or key output · GREY fill = group header · NAVY = section / year headers (A = built from actual data, E = estimate years).'),
 ('One assumption per cell', 'Every parameter lives once, on Inputs. Formula tabs never contain a typed number except the fixed 2024 cross-section on Calibration and Curves I:K.'),
 ('Check before you trust', 'Checks tab must read OK everywhere after Excel recalculates (it does on open). A FAIL means a reference broke in the copy.'),
 ('', None),
 ('WHERE THE METHOD IS WRITTEN UP', None),
 ('docs/AGE_CURVES.md', 'S, R, P: evidence, fit, validation, manual pulls.'), ('MODEL_BLUEPRINT.md', 'Architecture, mechanisms A–E, build order.'), ('findings.md', 'Lab notebook, Addenda 14–16.'), ('PROVENANCE.md §5', 'Every host and file touched.'),
 ('', None),
 ('NOT IN THIS WORKBOOK', None),
 ('Licensed material', 'S&P transcripts, the Stephens preview and the Black Diamond SOLS model are read locally only. The IHS-sourced ORNL tables (3.11 / 3.12 / 3.13, "further reproduction prohibited") are cited by table number, not reproduced.'),
 ('Lot-level rows', '1.3M sitemap rows live in data/cprt.db, not here.'),
]
for i, (a, b) in enumerate(readme):
    r = 5 + i
    if b is None and a: bar(wrm, r, 1, 2, a, BLUEF, WHITE_B, 16)
    elif a: put(wrm, f'A{r}', a, BOLD); put(wrm, f'B{r}', b, BLK, wrap=True); wrm.row_dimensions[r].height = max(15, 14 * (1 + len(b) // 105))
widths(wrm, {'A': 34, 'B': 120})

# ================================================================ Data tabs
for name, fname, title in DATA_TABS: add_data_sheet(wb, name, fname, title)

# ================================================================ Divider tabs (SOLS style) and sheet order
def divider(title):
    ws = wb.create_sheet(title); ws.sheet_view.showGridLines = False; ws.sheet_properties.tabColor = '0070C0'
    put(ws, 'B3', title.replace(' →', ''), font(C_NAVY, 20, bold=True)); put(ws, 'B5', 'section divider — nothing on this sheet', NOTE); ws.column_dimensions['B'].width = 60
    return ws
for d in ('Inputs & Curves →', 'Fleet Build →', 'Regressions →', 'Raw Data →'): divider(d)
ORDER = ['Cover', 'README', 'Inputs & Curves →', 'Inputs', 'Curves', 'Calibration', 'Fleet Build →', 'Data_Sales', 'Data_EPA_Survival', 'Survival', 'Fleet', 'FleetByAge', 'TLF_Roll',
         'Regressions →', 'Spread_Reg', 'RPU_Reg', 'Checks', 'Raw Data →', 'Data_CCC_Targets'] + [n for n, _, _ in DATA_TABS]
wb._sheets = [wb[n] for n in ORDER]

# ================================================================ Cover
wcv.sheet_view.showGridLines = False; wcv.sheet_properties.tabColor = '0070C0'
for c_, w in {'A': 3, 'B': 26, 'C': 30, 'D': 3, 'E': 30, 'F': 70}.items(): wcv.column_dimensions[c_].width = w
bar(wcv, 2, 2, 6, 'Copart, Inc. — intermediate build book', TITLEF, font(F_WHITE, 16, bold=True), 28)
put(wcv, 'B4', 'Ticker', BOLD); put(wcv, 'C4', 'NASDAQ: CPRT', font('674AE2', 14, bold=True))
put(wcv, 'B5', 'Current price', BLK); put(wcv, 'C5', None, BLUE, '"$"#,##0.00'); put(wcv, 'E5', '← type; blue = input', NOTE)
put(wcv, 'B6', 'Price target', BLK); put(wcv, 'C6', None, BLUE, '"$"#,##0.00'); put(wcv, 'E6', '← link to CPRT Model once built', NOTE)
put(wcv, 'B7', 'Implied return', BLK); put(wcv, 'C7', '=IF(OR(C5="",C6=""),"",C6/C5-1)', BLK, PCT2)
put(wcv, 'B9', 'Author', BOLD); put(wcv, 'C9', 'Kendall Wu', BLK); put(wcv, 'B10', 'Built', BOLD); put(wcv, 'C10', TODAY, BLK)
put(wcv, 'B11', 'Source of truth', BOLD); put(wcv, 'C11', 'data/csv + scripts/build_intermediate_xlsx.py (github kwu/cprt)', BLK)
rules(wcv, 4, 11, 2, 3)
bar(wcv, 13, 2, 6, 'Table of contents', NAVY, WHITE_B, 18)
header_row(wcv, 14, 2, 6, ['Tab', 'Section', '', '', 'What it holds'], wrap=False, center=False)
toc = [
 ('README', 'Read me', 'How to copy tabs, colour legend, where the method is written up.'),
 ('Inputs', 'Inputs & Curves', 'Every parameter once: survival stretch, R and P, exposure, future sales, fleet anchors; k(t) by year.'),
 ('Curves', 'Inputs & Curves', 'R(a), P(a), miles by age; the 7+ vs 0–6 miles / coverage decomposition.'),
 ('Calibration', 'Inputs & Curves', 'Solver set-up: five knobs → eight CCC 2024 statistics → one loss cell.'),
 ('Data_Sales', 'Fleet Build', 'Sales(MY): new-vehicle sales by model year, cars and light trucks, 1970–2030.'),
 ('Data_EPA_Survival', 'Fleet Build', 'S(age): EPA survival schedule by age (ORNL TEDB Ed.40 T3.15).'),
 ('Survival', 'Fleet Build', 'S(a,t): the EPA schedule stretched by k(t), cars and light trucks, ages 0–45 × 2013–2030.'),
 ('Fleet', 'Fleet Build', 'Vehicles on the road by model year × calendar year; total VIO.'),
 ('FleetByAge', 'Fleet Build', 'Same fleet re-indexed by age; VIO, share 7+/15+, count ≤6, average age.'),
 ('TLF_Roll', 'Fleet Build', 'Baseline total-loss frequency from demographics, its drift, and the out-of-sample checks vs CCC.'),
 ('Spread_Reg', 'Regressions', 'ΔTLF = a + b · spread(t−1): pairs built by formula, SLOPE / INTERCEPT / RSQ, prediction, live 2Q26 test.'),
 ('RPU_Reg', 'Regressions', 'ASP ~ used-car CPI by fiscal quarter; service and total RPU ~ ASP; the chain calculator.'),
 ('Checks', 'Regressions', 'OK / FAIL on the key outputs and the reference regression values.'),
 ('Data_CCC_Targets', 'Raw Data', 'The CCC statistics R and P are fitted to and tested against, with verbatim quotes and pages.'),
] + [(n, 'Raw Data', t) for n, _, t in DATA_TABS]
for i, (tabn, sec, desc) in enumerate(toc):
    r = 15 + i
    c = put(wcv, f'B{r}', tabn, font('0563C1', 11, underline='single')); c.hyperlink = f"#'{tabn}'!A1"
    put(wcv, f'C{r}', sec, BLK); put(wcv, f'F{r}', desc, BLK)
rules(wcv, 15, 14 + len(toc), 2, 3, header_row_=14); rules(wcv, 15, 14 + len(toc), 6, 6, header_row_=14)
r = 16 + len(toc)
bar(wcv, r, 2, 6, 'Colour legend', NAVY, WHITE_B, 18)
for i, (txt, f_, fl_) in enumerate([('Hard-coded input', BLUE, None), ('Formula', BLK, None), ('Link to another sheet', GRN, None), ('Calibrated parameter or key output', BOLD, PINK), ('Group header', BOLD, GREY), ('Section or year header  (A = actual data, E = estimate years)', WHITE_B, NAVY)]):
    put(wcv, f'B{r+1+i}', txt, f_, fl=fl_); put(wcv, f'C{r+1+i}', ('1,234.5' if i < 3 else ''), f_, fl=fl_)
rules(wcv, r + 1, r + 6, 2, 3)

OUT.parent.mkdir(exist_ok=True); wb.save(OUT)
nform = sum(1 for s in wb.worksheets for row in s.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('='))
print(f"-> {OUT.relative_to(ROOT)}  sheets={len(wb.worksheets)}  formulas={nform}")
