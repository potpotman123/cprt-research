"""Read-only inventory and source reconciliation of the user-supplied repository workbook."""
import csv,hashlib,json
from pathlib import Path
import openpyxl
HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[1]
P=Path('/Users/kwu/Library/Containers/com.apple.MobileSMS/Data/tmp/TemporaryItems/com.apple.MobileSMS/Media/6B7AB363-5EBE-4CCC-BA2F-CA9F98ABE187/CPRT_Intermediate (1).xlsx')
rows={int(r['year']):r for r in csv.DictReader(l for l in (ROOT/'data/csv/light_vehicle_sales_by_year.csv').read_text().splitlines() if not l.startswith('#'))}
w=openpyxl.load_workbook(P,read_only=True,data_only=True);inventory=[];hits=[];comparison=[]
for s in w:
    inventory.append({'sheet':s.title,'rows':s.max_row,'columns':s.max_column})
    for row in s:
        for cell in row:
            if isinstance(cell.value,str) and any(t in cell.value.lower() for t in ['pickup','crossover','minivan','car suv','truck suv','production share']):hits.append({'sheet':s.title,'cell':cell.coordinate,'text':cell.value})
for row in w['Data_Sales'].iter_rows(min_row=5,max_row=60):
    year=row[0].value
    if year not in rows:continue
    for idx,key in [(1,'cars_k'),(2,'light_trucks_k'),(3,'heavy_trucks_k')]:
        comparison.append({'year':year,'metric':key,'workbook_cell':row[idx].coordinate,'workbook_value':row[idx].value,'repo_value':float(rows[year][key]),'difference':row[idx].value-float(rows[year][key])})
w.close()
out={'path':str(P),'sha256':hashlib.sha256(P.read_bytes()).hexdigest(),'worksheets':inventory,'subtype_term_matches':hits,'comparison_count':len(comparison),'max_abs_sales_difference':max(abs(r['difference']) for r in comparison),'sales_comparison':comparison,'conclusion':'Workbook contains existing aggregate sales/age data; no subtype-history table identified. Formula results were not recalculated or treated as fresh model validation. Workbook not modified.','period_warning':'Data_Sales title says model year; source note describes calendar-year sales for FRED. Preserve actual source basis rather than inheriting the header.'}
(HERE/'workbook_inventory.json').write_text(json.dumps(out,indent=2)+'\n')
print(json.dumps({k:v for k,v in out.items() if k not in ['worksheets','sales_comparison']},indent=2))
