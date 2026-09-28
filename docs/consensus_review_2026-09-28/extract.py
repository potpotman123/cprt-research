"""Read-only extraction of supplied consensus; no workbook modification."""
import csv,hashlib,json,re
from pathlib import Path
import openpyxl
P=Path('/Users/kwu/Downloads/SPGlobal_Copart,Inc._Consensus_28-Sep-2026.xlsx')
OUT=Path(__file__).parent
w=openpyxl.load_workbook(P,data_only=True,read_only=True);s=w['Consensus']
def value(r,c):
 v=s.cell(r,c).value
 if v is None or v=='NA':return None
 if isinstance(v,(int,float)):return v
 return float(re.sub(r' [AE]$','',str(v)).replace(',','').replace('(','-').replace(')',''))
rows=[]
for c in range(10,18):
 rows.append({'period':s.cell(57,c).value,'total_revenue_mean_m':value(58,c),'total_revenue_median_m':value(61,c),'high_m':value(62,c),'low_m':value(63,c),'stddev_m':value(64,c),'analyst_count_label':s.cell(65,c).value,'prior_year_total_revenue_m':value(58,c-4),'yoy':value(58,c)/value(58,c-4)-1,'source_range':f'{openpyxl.utils.get_column_letter(c)}58:{openpyxl.utils.get_column_letter(c)}65'})
with (OUT/'total_revenue_consensus.csv').open('w') as f:
 wr=csv.DictWriter(f,fieldnames=list(rows[0]));wr.writeheader();wr.writerows(rows)
manifest={'source':str(P),'sha256':hashlib.sha256(P.read_bytes()).hexdigest(),'sheet':'Consensus','date_basis':'2026-09-28 filename; contributor-specific update dates absent','currency':'USD','scale':'millions','scope':'Consolidated total revenue; includes purchased vehicle sales','service_forecast_available':False,'price_or_target_available':False,'h1_fy27_revenue_m':sum(r['total_revenue_mean_m'] for r in rows[:2]),'h1_growth':sum(r['total_revenue_mean_m'] for r in rows[:2])/sum(r['prior_year_total_revenue_m'] for r in rows[:2])-1,'fy27_sum_quarters_m':sum(r['total_revenue_mean_m'] for r in rows[:4]),'fy27_growth':sum(r['total_revenue_mean_m'] for r in rows[:4])/sum(r['prior_year_total_revenue_m'] for r in rows[:4])-1,'note':'Sum of quarterly consensus means, not independently sourced annual consensus. Aggregated broker sets may differ across metrics/periods.'}
(OUT/'manifest.json').write_text(json.dumps(manifest,indent=2));print(json.dumps(manifest,indent=2))
